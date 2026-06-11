from __future__ import annotations

import hashlib
import mimetypes
import re
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import DocumentExtraction, SourceDocument, SourceDocumentChunk, SourceDocumentExtraction
from app.services.kyc_embeddings import LocalHashEmbeddingClient


TOKEN_WORD_RATIO = 1.35
SECTION_HEADING_PATTERN = re.compile(
    r"^\s*(?:\d+(?:\.\d+)*[\).:-]?\s*)?"
    r"(Scope of Work|Engagement Scope|Project Scope|Statement of Work|Timeframe and Payment Schedule|"
    r"Renewal Terms?|Commercial Terms?|Commercial Summary|Pricing|Fees|Billing Terms|Payment Terms|"
    r"Customer Responsibilities|Client Responsibilities|Stakeholders?|Deliverables?|Milestones?|"
    r"Risks?|Assumptions?|Dependencies?|Resources?|Service Lines?|Contract Terms?)\s*:?\s*$",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class PageText:
    page_number: int
    raw_text: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class ExtractedText:
    raw_text: str
    normalized_text: str
    page_count: int
    metadata: dict[str, Any]
    ocr_status: str | None = None
    ocr_engine: str | None = None
    page_texts: list[PageText] = field(default_factory=list)


class KycDocumentExtractionService:
    extractor_name = "local-document-extractor"
    extractor_version = "v1"

    def __init__(self, db: Session) -> None:
        self.db = db
        self.settings = get_settings()

    def extract_document(self, document: SourceDocument, *, force: bool = False) -> SourceDocumentExtraction:
        latest = self.latest_extraction(document.id)
        if latest and latest.status == "completed" and not force:
            if not self.page_extractions_for_document(document.id):
                self._store_page_extractions(
                    document,
                    latest,
                    ExtractedText(
                        raw_text=latest.raw_text or "",
                        normalized_text=latest.normalized_text or normalize_text(latest.raw_text or ""),
                        page_count=latest.page_count or 1,
                        metadata=latest.metadata_json or {},
                        page_texts=[PageText(page_number=1, raw_text=latest.raw_text or latest.normalized_text or "", metadata={"extractor": latest.extractor_name, "backfilled": True})],
                    ),
                    force=False,
                )
                self.db.flush()
            return latest

        now = datetime.now(timezone.utc)
        document.extraction_status = "running"
        document.extraction_started_at = now
        document.extraction_error = None
        path = self._document_path(document)
        extraction = SourceDocumentExtraction(
            source_document_id=document.id,
            status="running",
            extractor_name=self.extractor_name,
            extractor_version=self.extractor_version,
            mime_type=document.mime_type or self._guess_mime_type(document),
            started_at=now,
        )
        self.db.add(extraction)
        self.db.flush()

        if not self.settings.kyc_document_extraction_enabled:
            return self._finish_failure(document, extraction, "Document extraction is disabled.", status_value="failed")
        if path is None:
            return self._finish_failure(document, extraction, "No readable stored file path is available for this source document.", status_value="failed")
        if not path.exists():
            return self._finish_failure(document, extraction, f"Stored source document file was not found: {path.name}", status_value="failed")
        if path.stat().st_size > self.settings.kyc_document_extraction_max_file_mb * 1024 * 1024:
            return self._finish_failure(document, extraction, "Source document exceeds the configured extraction file-size limit.", status_value="failed")

        try:
            data = path.read_bytes()
            document.checksum_sha256 = hashlib.sha256(data).hexdigest()
            document.size_bytes = len(data)
            document.mime_type = document.mime_type or mimetypes.guess_type(path.name)[0] or "application/octet-stream"
            result = self._extract_by_type(path, document.mime_type)
        except Exception as exc:  # pragma: no cover - library boundary guard
            return self._finish_failure(document, extraction, f"Document extraction failed: {str(exc)[:400]}", status_value="failed")

        status_value = self._status_for_extracted_text(result)
        extraction.status = status_value
        extraction.raw_text = result.raw_text
        extraction.normalized_text = result.normalized_text
        extraction.page_count = result.page_count
        extraction.metadata_json = result.metadata
        extraction.error_message = self._extraction_error_for_status(result, status_value)
        extraction.completed_at = datetime.now(timezone.utc)
        document.pages = result.page_count or document.pages
        document.extracted_text_checksum = hashlib.sha256(result.normalized_text.encode("utf-8")).hexdigest() if result.normalized_text else None
        document.extraction_status = status_value
        document.extraction_completed_at = extraction.completed_at
        document.extraction_error = extraction.error_message
        document.ocr_status = result.ocr_status
        document.ocr_engine = result.ocr_engine
        self._store_page_extractions(document, extraction, result, force=force)
        self.db.flush()
        return extraction

    def latest_extraction(self, source_document_id: str) -> SourceDocumentExtraction | None:
        return self.db.scalar(
            select(SourceDocumentExtraction)
            .where(SourceDocumentExtraction.source_document_id == source_document_id)
            .order_by(SourceDocumentExtraction.created_at.desc())
            .limit(1)
        )

    def chunks_for_document(self, source_document_id: str) -> list[SourceDocumentChunk]:
        return list(
            self.db.scalars(
                select(SourceDocumentChunk)
                .where(SourceDocumentChunk.source_document_id == source_document_id)
                .order_by(SourceDocumentChunk.chunk_index)
            )
        )

    def page_extractions_for_document(self, source_document_id: str) -> list[DocumentExtraction]:
        return list(
            self.db.scalars(
                select(DocumentExtraction)
                .where(DocumentExtraction.document_id == source_document_id)
                .order_by(DocumentExtraction.page_number)
            )
        )

    def extract_stored_file(self, path: Path, mime_type: str | None) -> ExtractedText:
        if not path.exists():
            raise ValueError(f"Stored source document file was not found: {path.name}")
        if path.stat().st_size > self.settings.kyc_document_extraction_max_file_mb * 1024 * 1024:
            raise ValueError("Source document exceeds the configured extraction file-size limit.")
        return self._extract_by_type(path, mime_type or mimetypes.guess_type(path.name)[0])

    def ensure_chunks(self, document: SourceDocument, *, force: bool = False) -> list[SourceDocumentChunk]:
        existing = self.chunks_for_document(document.id)
        if existing and not force:
            return existing
        extraction = self.extract_document(document, force=force) if self._document_path(document) else self.latest_extraction(document.id)
        return self.chunk_document(document, extraction=extraction, force=True)

    def chunk_document(self, document: SourceDocument, *, extraction: SourceDocumentExtraction | None = None, force: bool = False) -> list[SourceDocumentChunk]:
        if force:
            self.db.execute(delete(SourceDocumentChunk).where(SourceDocumentChunk.source_document_id == document.id))
            self.db.flush()

        page_rows = self.page_extractions_for_document(document.id)
        chunks: list[dict[str, Any]] = []
        if page_rows:
            for row in page_rows:
                page_chunks = self._chunk_page_text(
                    row.raw_text,
                    row=row,
                    chunk_size_tokens=self.settings.ai_kyc_chunk_size_tokens,
                    overlap_tokens=self.settings.ai_kyc_chunk_overlap_tokens,
                )
                chunks.extend(page_chunks)
        else:
            text = (extraction.normalized_text if extraction else None) or ""
            chunks = self._chunk_text(
                text,
                chunk_size_tokens=self.settings.ai_kyc_chunk_size_tokens,
                overlap_tokens=self.settings.ai_kyc_chunk_overlap_tokens,
                section_label="Document Text",
                metadata={"chunk_strategy": "document_text"},
            )
        if not chunks:
            chunks = self._chunks_from_citations(document)
        embedding_client = self._local_embedding_client()
        vectors = embedding_client.embed([chunk["text"] for chunk in chunks]) if self.settings.ai_kyc_use_embeddings and chunks else []
        records: list[SourceDocumentChunk] = []
        for index, chunk in enumerate(chunks):
            chunk_hash = hashlib.sha256(chunk["text"].encode("utf-8")).hexdigest()
            vector = vectors[index] if index < len(vectors) else None
            record = SourceDocumentChunk(
                source_document_id=document.id,
                extraction_id=extraction.id if extraction else None,
                account_id=document.account_id,
                engagement_id=document.engagement_id,
                chunk_index=index,
                chunk_text=chunk["text"],
                chunk_hash=chunk_hash,
                page_number=chunk.get("page_number"),
                section_label=chunk.get("section_label"),
                start_offset=chunk.get("start_offset"),
                end_offset=chunk.get("end_offset"),
                token_count=estimate_tokens(chunk["text"]),
                sensitivity_level=self._chunk_sensitivity(document, chunk),
                source_type=document.source_type,
                trust_score=source_trust_score(document.source_type, is_sensitive=document.is_sensitive),
                embedding_provider=embedding_client.provider if vector else None,
                embedding_model=embedding_client.model if vector else None,
                embedding_json=vector,
                embedding_hash=chunk_hash if vector else None,
                metadata_json={
                    "source_title": document.title,
                    "source_file": document.file_name or document.title,
                    "chunk_strategy": "section_page_token",
                    **chunk.get("metadata", {}),
                },
            )
            self.db.add(record)
            records.append(record)
        self.db.flush()
        return records

    def _extract_by_type(self, path: Path, mime_type: str | None) -> ExtractedText:
        suffix = path.suffix.lower()
        if suffix == ".pdf" or mime_type == "application/pdf":
            return self._extract_pdf(path)
        if suffix == ".docx" or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return self._extract_docx(path)
        if suffix in {".txt", ".csv"} or (mime_type or "").startswith("text/"):
            return self._extract_text(path)
        if suffix in {".xlsx", ".xlsm"} or mime_type in {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        }:
            return self._extract_xlsx(path)
        if suffix == ".xls" or mime_type == "application/vnd.ms-excel":
            return self._extract_xls(path)
        raise ValueError(f"Unsupported source document type: {suffix or mime_type or 'unknown'}")

    def _extract_pdf(self, path: Path) -> ExtractedText:
        pages, extractor = self._extract_pdf_pages(path)
        raw = "\n\n".join(pages)
        normalized = normalize_text(raw)
        metadata = {
            "extractor_stack": extractor,
            "source_file": path.name,
            "pages": [{"page_number": index + 1, "char_count": len(value)} for index, value in enumerate(pages)],
        }
        page_texts = [PageText(page_number=index + 1, raw_text=value, metadata={"extractor": extractor}) for index, value in enumerate(pages)]
        if len(normalized) >= self.settings.kyc_ocr_min_text_chars or not self.settings.kyc_ocr_enabled:
            ocr_status = "not_required" if normalized else ("disabled" if not self.settings.kyc_ocr_enabled else "not_required")
            return ExtractedText(
                raw_text=raw,
                normalized_text=normalized,
                page_count=len(pages),
                metadata={**metadata, "ocr_status": ocr_status, "ocr_required": False},
                ocr_status=ocr_status,
                page_texts=page_texts,
            )
        ocr_pages = self._ocr_pdf_pages(path)
        if ocr_pages is None:
            return ExtractedText(
                raw_text=raw,
                normalized_text=normalized,
                page_count=len(pages),
                metadata={
                    **metadata,
                    "ocr_status": "ocr_required",
                    "ocr_required": True,
                    "ocr_attempted": True,
                    "ocr_engine": self.settings.kyc_ocr_engine,
                    "ocr_error": "OCR could not extract readable text. Install/configure OCR dependencies or retry after correcting the source file.",
                },
                ocr_status="ocr_required",
                ocr_engine=self.settings.kyc_ocr_engine,
                page_texts=page_texts,
            )
        ocr = "\n\n".join(ocr_pages)
        return ExtractedText(
            raw_text=ocr,
            normalized_text=normalize_text(ocr),
            page_count=len(ocr_pages),
            metadata={**metadata, "ocr_applied": True, "ocr_status": "completed", "ocr_required": True, "ocr_engine": self.settings.kyc_ocr_engine},
            ocr_status="completed",
            ocr_engine=self.settings.kyc_ocr_engine,
            page_texts=[PageText(page_number=index + 1, raw_text=value, metadata={"extractor": self.settings.kyc_ocr_engine}) for index, value in enumerate(ocr_pages)],
        )

    def _extract_pdf_pages(self, path: Path) -> tuple[list[str], str]:
        for extractor_name, extractor in (
            ("pymupdf", self._extract_pdf_pages_pymupdf),
            ("pdfplumber", self._extract_pdf_pages_pdfplumber),
            ("pypdf", self._extract_pdf_pages_pypdf),
        ):
            pages = extractor(path)
            if normalize_text("\n".join(pages)):
                return pages, extractor_name
        return [], "none"

    @staticmethod
    def _extract_pdf_pages_pymupdf(path: Path) -> list[str]:
        try:
            import fitz
        except Exception:
            return []
        try:
            with fitz.open(str(path)) as document:
                return [page.get_text("text") or "" for page in document]
        except Exception:
            return []

    @staticmethod
    def _extract_pdf_pages_pdfplumber(path: Path) -> list[str]:
        try:
            import pdfplumber
        except Exception:
            return []
        try:
            with pdfplumber.open(str(path)) as document:
                return [page.extract_text() or "" for page in document.pages]
        except Exception:
            return []

    @staticmethod
    def _extract_pdf_pages_pypdf(path: Path) -> list[str]:
        try:
            from pypdf import PdfReader
        except Exception:
            return []
        try:
            reader = PdfReader(str(path))
            return [page.extract_text() or "" for page in reader.pages]
        except Exception:
            return []

    def _ocr_pdf_pages(self, path: Path) -> list[str] | None:
        pages = self._ocr_pdf_pages_ocrmypdf(path)
        if pages and normalize_text("\n".join(pages)):
            return pages
        pages = self._ocr_pdf_pages_tesseract(path)
        return pages if pages and normalize_text("\n".join(pages)) else None

    def _ocr_pdf_pages_ocrmypdf(self, path: Path) -> list[str] | None:
        try:
            import ocrmypdf  # noqa: F401
        except Exception:
            return None
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                output = Path(tmpdir) / "ocr-output.pdf"
                subprocess.run(
                    ["ocrmypdf", "--skip-text", "--quiet", str(path), str(output)],
                    check=True,
                    timeout=max(60, self.settings.ai_kyc_timeout_seconds),
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
                pages, _ = self._extract_pdf_pages(output)
                return pages if normalize_text("\n".join(pages)) else None
        except Exception:
            return None

    def _ocr_pdf_pages_tesseract(self, path: Path) -> list[str] | None:
        try:
            from pdf2image import convert_from_path
            import pytesseract
        except Exception:
            return None
        try:
            images = convert_from_path(str(path), dpi=200)
            return [pytesseract.image_to_string(image) for image in images]
        except Exception:
            return None

    @staticmethod
    def _extract_docx(path: Path) -> ExtractedText:
        from docx import Document

        document = Document(str(path))
        parts: list[str] = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    parts.append(" | ".join(values))
        raw = "\n".join(parts)
        return ExtractedText(
            raw_text=raw,
            normalized_text=normalize_text(raw),
            page_count=1 if raw.strip() else 0,
            metadata={"paragraphs": len(parts), "pages": [{"page_number": 1, "char_count": len(raw)}]},
            ocr_status="not_required",
            page_texts=[PageText(page_number=1, raw_text=raw, metadata={"extractor": "python-docx"})] if raw.strip() else [],
        )

    @staticmethod
    def _extract_text(path: Path) -> ExtractedText:
        raw = path.read_text(encoding="utf-8", errors="ignore")
        return ExtractedText(
            raw_text=raw,
            normalized_text=normalize_text(raw),
            page_count=1 if raw.strip() else 0,
            metadata={"line_count": raw.count("\n") + 1, "pages": [{"page_number": 1, "char_count": len(raw)}]},
            ocr_status="not_required",
            page_texts=[PageText(page_number=1, raw_text=raw, metadata={"extractor": "plain-text"})] if raw.strip() else [],
        )

    @staticmethod
    def _extract_xlsx(path: Path) -> ExtractedText:
        from openpyxl import load_workbook

        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        sheet_parts: list[str] = []
        page_texts: list[PageText] = []
        pages_metadata: list[dict[str, Any]] = []
        for index, sheet in enumerate(workbook.worksheets, start=1):
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    rows.append(" | ".join(values))
            sheet_text = "\n".join(rows)
            if not sheet_text.strip():
                continue
            sheet_parts.append(f"[Sheet: {sheet.title}]\n{sheet_text}")
            page_texts.append(
                PageText(
                    page_number=index,
                    raw_text=sheet_text,
                    metadata={"extractor": "openpyxl", "sheet_name": sheet.title},
                )
            )
            pages_metadata.append({"page_number": index, "sheet_name": sheet.title, "char_count": len(sheet_text)})

        raw = "\n\n".join(sheet_parts)
        return ExtractedText(
            raw_text=raw,
            normalized_text=normalize_text(raw),
            page_count=len(page_texts),
            metadata={"extractor": "openpyxl", "sheet_count": len(workbook.worksheets), "pages": pages_metadata},
            ocr_status="not_required",
            page_texts=page_texts,
        )

    @staticmethod
    def _extract_xls(path: Path) -> ExtractedText:
        import xlrd

        workbook = xlrd.open_workbook(str(path))
        sheet_parts: list[str] = []
        page_texts: list[PageText] = []
        pages_metadata: list[dict[str, Any]] = []
        for index, sheet in enumerate(workbook.sheets(), start=1):
            rows: list[str] = []
            for row_index in range(sheet.nrows):
                values = [
                    str(sheet.cell_value(row_index, column_index)).strip()
                    for column_index in range(sheet.ncols)
                    if str(sheet.cell_value(row_index, column_index)).strip()
                ]
                if values:
                    rows.append(" | ".join(values))
            sheet_text = "\n".join(rows)
            if not sheet_text.strip():
                continue
            sheet_parts.append(f"[Sheet: {sheet.name}]\n{sheet_text}")
            page_texts.append(
                PageText(
                    page_number=index,
                    raw_text=sheet_text,
                    metadata={"extractor": "xlrd", "sheet_name": sheet.name},
                )
            )
            pages_metadata.append({"page_number": index, "sheet_name": sheet.name, "char_count": len(sheet_text)})

        raw = "\n\n".join(sheet_parts)
        return ExtractedText(
            raw_text=raw,
            normalized_text=normalize_text(raw),
            page_count=len(page_texts),
            metadata={"extractor": "xlrd", "sheet_count": workbook.nsheets, "pages": pages_metadata},
            ocr_status="not_required",
            page_texts=page_texts,
        )

    @staticmethod
    def _status_for_extracted_text(result: ExtractedText) -> str:
        if result.ocr_status == "ocr_required":
            return "ocr_required"
        if not result.normalized_text.strip():
            return "needs_review"
        return "completed"

    @staticmethod
    def _extraction_error_for_status(result: ExtractedText, status_value: str) -> str | None:
        if status_value == "ocr_required":
            return "OCR is required before readable text can be extracted from this source document."
        if status_value == "needs_review":
            return "No readable text was extracted from this source document."
        return None

    def _store_page_extractions(self, document: SourceDocument, extraction: SourceDocumentExtraction, result: ExtractedText, *, force: bool) -> None:
        if force:
            self.db.execute(delete(DocumentExtraction).where(DocumentExtraction.document_id == document.id))
            self.db.flush()
        page_texts = result.page_texts or ([PageText(page_number=1, raw_text=result.raw_text)] if result.raw_text.strip() else [])
        for page in page_texts:
            raw_text = page.raw_text or ""
            if not raw_text.strip():
                continue
            checksum = hashlib.sha256(raw_text.encode("utf-8")).hexdigest()
            self.db.add(
                DocumentExtraction(
                    document_id=document.id,
                    extraction_id=extraction.id,
                    raw_text=raw_text,
                    page_number=page.page_number,
                    source_file=document.file_name or document.title,
                    checksum=checksum,
                    extractor_name=self.extractor_name,
                    extractor_version=self.extractor_version,
                    metadata_json=page.metadata,
                )
            )

    def _finish_failure(self, document: SourceDocument, extraction: SourceDocumentExtraction, message: str, *, status_value: str) -> SourceDocumentExtraction:
        now = datetime.now(timezone.utc)
        extraction.status = status_value
        extraction.error_message = message
        extraction.completed_at = now
        document.extraction_status = status_value
        document.extraction_completed_at = now
        document.extraction_error = message
        self.db.flush()
        return extraction

    def _document_path(self, document: SourceDocument) -> Path | None:
        candidates = [document.storage_path, document.file_url]
        for candidate in candidates:
            if not candidate or str(candidate).startswith(("http://", "https://")):
                continue
            path = Path(candidate)
            if path.exists() or path.is_absolute():
                return path
            local_path = Path(self.settings.local_content_storage_dir) / candidate
            if local_path.exists():
                return local_path
        if document.file_name:
            local_path = Path(self.settings.local_content_storage_dir) / document.file_name
            if local_path.exists():
                return local_path
        return None

    @staticmethod
    def _guess_mime_type(document: SourceDocument) -> str | None:
        return mimetypes.guess_type(document.file_name or document.file_url or document.storage_path or "")[0]

    def _chunk_page_text(
        self,
        text: str,
        *,
        row: DocumentExtraction,
        chunk_size_tokens: int,
        overlap_tokens: int,
    ) -> list[dict[str, Any]]:
        sections = self._split_sections(text)
        chunks: list[dict[str, Any]] = []
        for section_index, section in enumerate(sections):
            section_chunks = self._chunk_text(
                normalize_text(section["text"]),
                chunk_size_tokens=chunk_size_tokens,
                overlap_tokens=overlap_tokens,
                section_label=section["section_label"],
                metadata={
                    "document_extraction_id": row.id,
                    "page_checksum": row.checksum,
                    "source_file": row.source_file,
                    "section_index": section_index,
                    "section_start_line": section.get("start_line"),
                    "section_end_line": section.get("end_line"),
                    "section_heading_raw": section.get("raw_heading"),
                },
            )
            for chunk in section_chunks:
                chunks.append({**chunk, "page_number": row.page_number})
        return chunks

    @staticmethod
    def _chunk_text(
        text: str,
        *,
        chunk_size_tokens: int,
        overlap_tokens: int,
        section_label: str | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> list[dict[str, Any]]:
        words = text.split()
        if not words:
            return []
        words_per_chunk = max(80, round(chunk_size_tokens / TOKEN_WORD_RATIO))
        overlap_words = max(0, min(round(overlap_tokens / TOKEN_WORD_RATIO), words_per_chunk // 2))
        chunks: list[dict[str, Any]] = []
        start = 0
        while start < len(words):
            end = min(len(words), start + words_per_chunk)
            chunk_text = " ".join(words[start:end]).strip()
            if chunk_text:
                chunks.append(
                    {
                        "text": chunk_text,
                        "start_offset": start,
                        "end_offset": end,
                        "section_label": section_label,
                        "metadata": dict(metadata or {}),
                    }
                )
            if end >= len(words):
                break
            start = max(end - overlap_words, start + 1)
        return chunks

    def _split_sections(self, text: str) -> list[dict[str, Any]]:
        lines = text.splitlines()
        sections: list[dict[str, Any]] = []
        current_heading = "Document Text"
        current_raw_heading: str | None = None
        current_lines: list[str] = []
        current_start_line = 1
        saw_heading = False

        def flush(end_line: int) -> None:
            body = "\n".join(current_lines).strip()
            if not body:
                return
            sections.append(
                {
                    "section_label": current_heading,
                    "raw_heading": current_raw_heading,
                    "start_line": current_start_line,
                    "end_line": end_line,
                    "text": body,
                }
            )

        for line_number, line in enumerate(lines, start=1):
            heading = self._section_label_from_line(line)
            if heading:
                flush(line_number - 1)
                saw_heading = True
                current_heading = heading
                current_raw_heading = line.strip()
                current_lines = [line.strip()]
                current_start_line = line_number
            else:
                current_lines.append(line)
        flush(len(lines))

        if not sections:
            normalized = normalize_text(text)
            return [{"section_label": "Document Text", "raw_heading": None, "start_line": 1, "end_line": len(lines) or 1, "text": normalized}] if normalized else []
        if not saw_heading and len(sections) == 1:
            sections[0]["section_label"] = "Document Text"
        return sections

    @staticmethod
    def _section_label_from_line(line: str) -> str | None:
        cleaned = re.sub(r"\s+", " ", line or "").strip()
        if not cleaned or len(cleaned) > 140:
            return None
        match = SECTION_HEADING_PATTERN.match(cleaned)
        if not match:
            return None
        lower = cleaned.lower()
        if any(token in lower for token in ("renewal", "notice")):
            return "Renewal Terms"
        if any(token in lower for token in ("commercial", "pricing", "fees", "billing", "payment")):
            return "Commercial Terms"
        if any(token in lower for token in ("timeframe", "schedule", "contract", "term")):
            return "Contract Terms"
        if any(token in lower for token in ("scope", "statement of work", "service", "resources")):
            return "Scope of Work"
        if any(token in lower for token in ("deliverable", "milestone")):
            return "Deliverables"
        if any(token in lower for token in ("risk", "assumption", "dependenc")):
            return "Risks And Assumptions"
        if any(token in lower for token in ("stakeholder", "responsibilit", "customer", "client")):
            return "Stakeholders And Responsibilities"
        return cleaned[:120]

    @staticmethod
    def _chunk_sensitivity(document: SourceDocument, chunk: dict[str, Any]) -> str:
        if document.is_sensitive:
            return "sensitive"
        text = f"{chunk.get('section_label') or ''} {chunk.get('text') or ''}".lower()
        if document.source_type == "commercial_note" or any(token in text for token in ("commercial", "pricing", "fee", "billing", "payment", "contract value", "renewal", "notice period")):
            return "commercial"
        return "standard"

    def _local_embedding_client(self) -> LocalHashEmbeddingClient:
        model = self.settings.ai_kyc_embedding_model if self.settings.ai_kyc_embedding_provider == "local_hash" else "local-hash-v1"
        return LocalHashEmbeddingClient(dimensions=self.settings.ai_kyc_embedding_dimensions, model=model or "local-hash-v1")

    @staticmethod
    def _chunks_from_citations(document: SourceDocument) -> list[dict[str, Any]]:
        chunks = []
        for citation in document.citations:
            if not citation.excerpt:
                continue
            chunks.append(
                {
                    "text": citation.excerpt,
                    "page_number": citation.page_number,
                    "section_label": citation.label,
                    "metadata": {"field_key": citation.field_key, "citation_id": citation.id, "chunk_strategy": "citation_fallback"},
                }
            )
        return chunks


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def estimate_tokens(value: str) -> int:
    return max(1, round(len((value or "").split()) * TOKEN_WORD_RATIO))


def source_trust_score(source_type: str, *, is_sensitive: bool = False) -> int:
    mapping = {
        "sow": 100,
        "contract": 100,
        "msa": 100,
        "project_charter": 92,
        "engagement": 95,
        "commercial_note": 90,
        "governance": 80,
        "timeline": 80,
        "account_note": 70,
        "fathom": 70,
        "research": 65,
        "prior_kyc": 50,
    }
    score = mapping.get(source_type, 60)
    return min(100, score + 3) if is_sensitive and source_type in {"sow", "contract", "commercial_note"} else score
