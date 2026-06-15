from collections.abc import Generator
from datetime import datetime, timedelta, timezone
from io import BytesIO
import json
import re

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base, get_db
from app.config import get_settings
from app.main import app
from app.models import Account, AccountOwner, CustomFieldValue, DocumentExtraction, Engagement, EngagementImportDraft, OnboardingDraft, Permission, Role, RolePermission, SourceDocument, SourceDocumentExtraction, Stakeholder
from app.services.kyc_document_extraction import KycDocumentExtractionService
from app.services.source_document_contract import SERVICE_LINE_LABELS, infer_service_lines_from_text
from app.services.sow_extraction import SowExtractionService
from app.services.seed import seed_default_data


@pytest.fixture()
def db_session() -> Generator[Session, None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(bind=engine)

    with TestingSessionLocal() as session:
        seed_default_data(session)
        yield session

    Base.metadata.drop_all(bind=engine)


@pytest.fixture()
def client(db_session: Session) -> Generator[TestClient, None, None]:
    def override_get_db() -> Generator[Session, None, None]:
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    test_client = TestClient(app)
    try:
        yield test_client
    finally:
        test_client.close()
        app.dependency_overrides.clear()


def auth_headers(client: TestClient, email: str = "admin@tkxel.com", password: str = "Admin@12345") -> dict[str, str]:
    response = client.post("/api/auth/login", json={"email": email, "password": password})
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def seeded_user(client: TestClient, headers: dict[str, str], role: str) -> dict:
    response = client.get("/api/admin/users", headers=headers, params={"role": role, "page": 1, "page_size": 1})
    assert response.status_code == 200
    return response.json()["items"][0]


def create_account_manager_user(client: TestClient, headers: dict[str, str], email: str, full_name: str) -> dict:
    response = client.post(
        "/api/admin/users",
        headers=headers,
        json={
            "email": email,
            "password": "User@12345",
            "full_name": full_name,
            "role": "account_manager",
            "title": "Account Manager",
            "avatar_initials": "".join(part[0] for part in full_name.split()[:2]).upper(),
            "is_active": True,
        },
    )
    assert response.status_code == 201
    return response.json()


def draft_payload(account_name: str, owner_id: str, *, include_engagement: bool = True) -> dict:
    account_slug = re.sub(r"[^a-z0-9]+", "-", account_name.lower()).strip("-") or "customer"
    engagement_drafts = []
    if include_engagement:
        engagement_drafts.append(
            {
                "name": "Customer intelligence modernization",
                "owner_id": owner_id,
                "service_lines": ["Account onboarding"],
                "value": 125000,
                "currency": "USD",
                "delivery_status": "active",
                "start_date": datetime.now(timezone.utc).replace(hour=12, minute=0, second=0, microsecond=0).isoformat().replace("+00:00", "Z"),
                "commercial_context": "Project Charter p1: account and engagement scope.",
                "risks": ["KYC has not been completed yet"],
                "source_citation": "Project Charter p1: account and engagement scope.",
                "confidence": 82,
            }
        )
    return {
        "account_name": account_name,
        "project_name": "Customer intelligence modernization",
        "company_url": f"https://{account_slug}.customer.example.com",
        "linkedin_url": "https://www.linkedin.com/company/customer-example",
        "lifecycle_status": "Onboarding",
        "segment": "Growth",
        "region": "Global",
        "commercial_value": 125000,
        "currency": "USD",
        "primary_owner_id": owner_id,
        "source_citation": "Project Charter p1: account and engagement scope.",
        "source_documents": [
            {
                "title": "Customer Project Charter",
                "source_type": "project_charter",
                "file_name": "customer-charter.pdf",
                "confidence": 91,
                "pages": 4,
                "citations": [
                    {
                        "label": "Project Charter p1",
                        "page_number": 1,
                        "excerpt": "Customer intelligence modernization scope.",
                        "field_key": "account_name",
                    }
                ],
            }
        ],
        "engagement_drafts": engagement_drafts,
    }


def iso_days_from_now(days: int) -> str:
    value = datetime.now(timezone.utc) + timedelta(days=days)
    return value.replace(hour=12, minute=0, second=0, microsecond=0).isoformat().replace("+00:00", "Z")


def date_days_from_now(days: int) -> str:
    return (datetime.now(timezone.utc).date() + timedelta(days=days)).isoformat()


def test_source_document_service_line_contract_is_centralized() -> None:
    assert "API integration" in SERVICE_LINE_LABELS
    service_lines = infer_service_lines_from_text("The team will handle React, API integration, cloud, and QA automation.")
    assert service_lines == ["React engineering", "QA automation", "Cloud integration", "API integration"]


def engagement_payload(owner_id: str, **overrides) -> dict:
    payload = {
        "name": "Strategic SOW QA",
        "description": "QA coverage for manually managed Engagement/SOW records.",
        "owner_id": owner_id,
        "service_lines": ["Engineering", "Customer Success"],
        "source_links": [{"title": "Signed SOW", "url": "https://customer.example.com/sow"}],
        "contract_value": 87500,
        "currency": "USD",
        "delivery_status": "active",
        "commercial_status": "healthy",
        "health_score": 82,
        "health_status": "green",
        "renewal_risk": "medium",
        "start_date": iso_days_from_now(-15),
        "end_date": iso_days_from_now(120),
        "renewal_date": iso_days_from_now(90),
        "notice_period_days": 45,
        "notice_deadline": iso_days_from_now(1),
        "resource_dependency_notes": "Named platform engineer availability.",
        "risks": ["Dependency on customer data access"],
        "source_citation": "SOW section 3.",
    }
    payload.update(overrides)
    return payload


ENGAGEMENT_RESPONSE_FIELDS = {
    "id",
    "account_id",
    "name",
    "description",
    "status",
    "owner_id",
    "owner_name",
    "service_lines",
    "source_document_ids",
    "source_links",
    "value",
    "contract_value",
    "currency",
    "delivery_status",
    "commercial_status",
    "delivery_health",
    "health_score",
    "health_status",
    "renewal_risk",
    "start_date",
    "end_date",
    "renewal_date",
    "notice_deadline",
    "notice_period_days",
    "days_to_expiry",
    "renewal_status",
    "resource_dependency_notes",
    "risks",
    "created_at",
    "updated_at",
    "created_by",
    "updated_by",
}


def create_engagement_for_account(client: TestClient, headers: dict[str, str], account_id: str, owner_id: str, **overrides) -> dict:
    response = client.post(f"/api/accounts/{account_id}/engagements", headers=headers, json=engagement_payload(owner_id, **overrides))
    assert response.status_code == 201
    return response.json()


def create_approved_account(client: TestClient, headers: dict[str, str], account_name: str, *, include_engagement: bool = True) -> tuple[str, str, dict]:
    owner = seeded_user(client, headers, "account_manager")
    draft_response = client.post("/api/onboarding/drafts", headers=headers, json=draft_payload(account_name, owner["id"], include_engagement=include_engagement))
    assert draft_response.status_code == 201

    approve_response = client.post(f"/api/onboarding/drafts/{draft_response.json()['id']}/approve", headers=headers)
    assert approve_response.status_code == 200
    approved = approve_response.json()
    return approved["approved_account_id"], owner["id"], approved


def create_account_without_engagement(db_session: Session, owner: dict, account_name: str) -> str:
    account = Account(
        name=account_name,
        project_name="Legacy account without engagement",
        company_url="https://legacy.example.com",
        linkedin_url="https://www.linkedin.com/company/legacy-account",
        lifecycle_status="Onboarding",
        segment="Growth",
        region="Global",
        risk_status="warning",
        created_by_id=owner["id"],
    )
    db_session.add(account)
    db_session.flush()
    db_session.add(
        AccountOwner(
            account_id=account.id,
            user_id=owner["id"],
            user_name=owner["full_name"],
            user_email=owner["email"],
            ownership_role="primary_am",
            is_primary=True,
            created_by_id=owner["id"],
        )
    )
    db_session.commit()
    return account.id


def test_onboarding_draft_approval_creates_account_sources_and_engagement(client: TestClient) -> None:
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")

    create_response = client.post("/api/onboarding/drafts", headers=headers, json=draft_payload("Northwind Workspace", owner["id"]))
    assert create_response.status_code == 201
    draft = create_response.json()
    assert draft["status"] == "ready_for_review"
    assert draft["linkedin_url"] == "https://www.linkedin.com/company/customer-example"
    assert draft["source_documents"][0]["citations"][0]["field_key"] == "account_name"
    assert draft["engagement_drafts"][0]["name"] == "Customer intelligence modernization"

    list_response = client.get("/api/onboarding/drafts", headers=headers, params={"search": "Northwind", "status": "ready_for_review", "page": 1, "page_size": 5})
    assert list_response.status_code == 200
    assert list_response.json()["total"] == 1

    approve_response = client.post(f"/api/onboarding/drafts/{draft['id']}/approve", headers=headers)
    assert approve_response.status_code == 200
    account_id = approve_response.json()["approved_account_id"]

    overview_response = client.get(f"/api/accounts/{account_id}/overview", headers=headers)
    assert overview_response.status_code == 200
    overview = overview_response.json()
    assert overview["account"]["name"] == "Northwind Workspace"
    assert overview["account"]["linkedin_url"] == "https://www.linkedin.com/company/customer-example"
    assert overview["account"]["lifecycle_status"] == "Active"
    assert overview["account"]["primary_owner"]["user_id"] == owner["id"]
    assert overview["account"]["has_health_score"] is False
    assert overview["account"]["health"] == {"overall": 0, "relationship": 0, "usage": 0, "delivery": 0, "commercial": 0}
    assert overview["engagements"]["total"] == 1
    assert overview["engagements"]["items"][0]["name"] == "Customer intelligence modernization"
    assert overview["engagements"]["items"][0]["status"] == "active"
    assert overview["attachments"]["total"] == 1

    rollup_response = client.get(f"/api/accounts/{account_id}/health/rollup", headers=headers)
    assert rollup_response.status_code == 200
    rollup = rollup_response.json()
    assert rollup["metric_version"] == "engagement-health-rollup-adapter-v1"
    assert rollup["overall"] == 0
    assert rollup["contributions"][0]["name"] == "Customer intelligence modernization"
    assert rollup["contributions"][0]["score"] == 82
    assert rollup["contributions"][0]["scoring_status"] == "pending_scoring_engine"


def test_onboarding_draft_update_edits_engagement_baseline_before_approval(client: TestClient) -> None:
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")

    create_response = client.post("/api/onboarding/drafts", headers=headers, json=draft_payload("Editable Baseline Account", owner["id"]))
    assert create_response.status_code == 201
    draft = create_response.json()
    engagement_id = draft["engagement_drafts"][0]["id"]
    end_date = iso_days_from_now(180)
    notice_deadline = iso_days_from_now(120)

    update_response = client.patch(
        f"/api/onboarding/drafts/{draft['id']}",
        headers=headers,
        json={
            "engagement_drafts": [
                {
                    "id": engagement_id,
                    "name": "Reviewed project validation baseline",
                    "value": 225000,
                    "end_date": end_date,
                    "renewal_date": end_date,
                    "notice_deadline": notice_deadline,
                    "auto_renewal": True,
                    "confidence": 91,
                    "ops_lead_name": "Delivery Lead Review",
                    "source_citation": "Reviewed baseline from SOW p2.",
                }
            ]
        },
    )
    assert update_response.status_code == 200
    updated_engagement = update_response.json()["engagement_drafts"][0]
    assert updated_engagement["name"] == "Reviewed project validation baseline"
    assert updated_engagement["value"] == 225000
    assert updated_engagement["notice_deadline"].startswith(notice_deadline[:10])
    assert updated_engagement["auto_renewal"] is True
    assert updated_engagement["confidence"] == 91
    assert updated_engagement["ops_lead_name"] == "Delivery Lead Review"
    assert updated_engagement["source_citation"] == "Reviewed baseline from SOW p2."

    approve_response = client.post(f"/api/onboarding/drafts/{draft['id']}/approve", headers=headers)
    assert approve_response.status_code == 200
    account_id = approve_response.json()["approved_account_id"]
    engagement_response = client.get(f"/api/accounts/{account_id}/engagements", headers=headers)
    assert engagement_response.status_code == 200
    engagement = engagement_response.json()["items"][0]
    assert engagement["name"] == "Reviewed project validation baseline"
    assert engagement["value"] == 225000
    assert engagement["end_date"].startswith(end_date[:10])
    assert engagement["notice_deadline"].startswith(notice_deadline[:10])
    assert engagement["auto_renewal"] is True
    assert engagement["delivery_health"] == 91
    assert engagement["ops_lead_name"] == "Delivery Lead Review"
    assert engagement["source_citation"] == "Reviewed baseline from SOW p2."


def test_onboarding_upload_extracts_draft_from_content_not_filename(client: TestClient) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")
    other_owner = create_account_manager_user(client, headers, "unrelated.source.manager@tkxel.com", "Unrelated Source Manager")
    content = b"""
Demo document only. Not a real signed commercial agreement.
Account Name: McDonald's Corporation
Company URL: https://www.mcdonalds.com
LinkedIn URL: https://www.linkedin.com/company/mcdonalds-corporation
Project Name: Digital Experience Modernization Program
Segment: Enterprise
Region: North America
Service Lines: Product Engineering, Data Engineering, QA Automation, Cloud Integration
Contract Value: USD 1250000
Start Date: 2026-07-01
End Date: 2027-06-30
Renewal Date: 2027-06-30
Notice Period: 60 days
Scope of Work:
Tkxel will support mobile ordering, loyalty personalization, restaurant operations dashboards, API integration, and release-quality automation.
Commercial Summary:
Fixed monthly delivery pod with milestone acceptance, monthly invoicing, and executive governance checkpoints.
Renewal Terms:
This SOW shall stand renewed automatically unless either Party gives 60 days written notice.
Risks:
Customer data access, franchise operating model complexity, and point-of-sale integration dependency.
"""

    try:
        response = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            data={"manager_id": owner["id"], "manager_name": owner["full_name"], "manager_email": owner["email"]},
            files=[("files", ("wrong-client-name.txt", content, "text/plain"))],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert response.status_code == 201
    draft = response.json()
    assert "McDonald's Corporation" in draft["source_documents"][0]["extracted_text"]
    assert "This SOW shall stand renewed automatically" in draft["source_documents"][0]["extracted_text"]
    assert draft["account_name"] == "McDonald's Corporation"
    assert draft["account_name"] != "Wrong Client Name"
    assert draft["project_name"] == "Digital Experience Modernization Program"
    assert draft["company_url"] == "https://www.mcdonalds.com"
    assert draft["linkedin_url"] == "https://www.linkedin.com/company/mcdonalds-corporation"
    assert draft["source_documents"][0]["file_name"] == "wrong-client-name.txt"
    assert draft["source_documents"][0]["extraction_status"] == "completed"
    assert draft["source_citation"].startswith("wrong-client-name.txt p1: Account name inferred")
    assert "McDonald's Corporation" in draft["source_citation"]
    field_keys = {citation["field_key"]: citation for citation in draft["source_documents"][0]["citations"]}
    assert field_keys["account_name"]["page_number"] == 1
    assert "McDonald's Corporation" in field_keys["account_name"]["excerpt"]
    assert field_keys["commercial_value"]["page_number"] == 1
    assert field_keys["start_date"]["page_number"] == 1

    download = client.get(
        f"/api/onboarding/drafts/{draft['id']}/documents/{draft['source_documents'][0]['id']}/download",
        headers=headers,
    )
    assert download.status_code == 200
    assert b"McDonald's Corporation" in download.content

    unrelated_download = client.get(
        f"/api/onboarding/drafts/{draft['id']}/documents/{draft['source_documents'][0]['id']}/download",
        headers=auth_headers(client, other_owner["email"], "User@12345"),
    )
    assert unrelated_download.status_code == 403

    approve_response = client.post(f"/api/onboarding/drafts/{draft['id']}/approve", headers=headers)
    assert approve_response.status_code == 200
    account_id = approve_response.json()["approved_account_id"]
    account_download = client.get(
        f"/api/accounts/{account_id}/attachments/{draft['source_documents'][0]['id']}/download",
        headers=headers,
    )
    assert account_download.status_code == 200
    assert b"Digital Experience Modernization Program" in account_download.content


def test_onboarding_upload_defaults_missing_engagement_start_date_for_one_click_approval(client: TestClient, db_session: Session) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")
    content = b"""
Statement of Work
Client Legal Name: Easy Approval Industries
Company URL: https://easy-approval.example.com
LinkedIn URL: https://www.linkedin.com/company/easy-approval-industries
Program Name: Customer Portal Support
Segment: Enterprise
Region: North America
Service Lines: Product Engineering, QA Automation
Scope of Work:
Tkxel will support the customer portal team with engineering delivery, QA automation, release planning, and operational governance.
Commercial Summary:
Monthly blended team model with account governance checkpoints.
"""

    try:
        response = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            data={"manager_id": owner["id"], "manager_name": owner["full_name"], "manager_email": owner["email"]},
            files=[("files", ("easy-approval-sow.txt", content, "text/plain"))],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert response.status_code == 201
    draft = response.json()
    assert draft["engagement_drafts"][0]["start_date"].startswith(datetime.now(timezone.utc).date().isoformat())

    draft_model = db_session.query(OnboardingDraft).filter(OnboardingDraft.id == draft["id"]).one()
    draft_model.engagement_drafts[0].start_date = None
    db_session.commit()

    approve_response = client.post(f"/api/onboarding/drafts/{draft['id']}/approve", headers=headers)
    assert approve_response.status_code == 200
    account_id = approve_response.json()["approved_account_id"]
    created_engagement = db_session.query(Engagement).filter(Engagement.account_id == account_id).one()
    assert created_engagement.start_date is not None
    assert created_engagement.start_date.date().isoformat() == datetime.now(timezone.utc).date().isoformat()


def test_onboarding_upload_uses_submitted_form_fields_as_account_source_of_truth(client: TestClient) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")
    content = b"""
Account Name: Extracted Document Customer
Company URL: https://wrong-document-customer.example.com
LinkedIn URL: https://www.linkedin.com/company/wrong-document-customer
Project Name: Extracted Document Project
Service Lines: Product Engineering, Cloud Integration
Contract Value: USD 250000
Start Date: 2026-07-01
End Date: 2027-06-30
Scope of Work:
The uploaded source should be preserved for KYC, but reviewed form fields should define account information.
"""

    try:
        response = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            data={
                "account_name": "Submitted Cafe Zupas",
                "project_name": "Submitted Guest Experience Program",
                "company_url": "https://submitted-cafe-zupas.example.com",
                "linkedin_url": "https://www.linkedin.com/company/submitted-cafe-zupas",
                "manager_id": owner["id"],
                "manager_name": owner["full_name"],
                "manager_email": owner["email"],
            },
            files=[("files", ("wrong-client-name.txt", content, "text/plain"))],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert response.status_code == 201
    draft = response.json()
    assert draft["account_name"] == "Submitted Cafe Zupas"
    assert draft["project_name"] == "Submitted Guest Experience Program"
    assert draft["company_url"] == "https://submitted-cafe-zupas.example.com"
    assert draft["linkedin_url"] == "https://www.linkedin.com/company/submitted-cafe-zupas"
    assert draft["source_citation"].endswith("account fields submitted in the creation form.")
    assert draft["engagement_drafts"][0]["name"] == "Submitted Guest Experience Program"
    assert "Extracted Document Customer" in draft["source_documents"][0]["extracted_text"]

    approve_response = client.post(f"/api/onboarding/drafts/{draft['id']}/approve", headers=headers)
    assert approve_response.status_code == 200
    account_id = approve_response.json()["approved_account_id"]
    overview_response = client.get(f"/api/accounts/{account_id}/overview", headers=headers)
    assert overview_response.status_code == 200
    assert overview_response.json()["account"]["name"] == "Submitted Cafe Zupas"
    assert overview_response.json()["account"]["project_name"] == "Submitted Guest Experience Program"


def test_onboarding_upload_extract_endpoint_prefills_without_creating_draft(client: TestClient, db_session: Session) -> None:
    headers = auth_headers(client)
    content = b"""
Account Name: Cafe Zupas
Project Name: Guest Experience Modernization
Company URL: https://www.cafezupas.com
LinkedIn URL: https://www.linkedin.com/company/cafe-zupas
Service Lines: Product Engineering, Cloud Integration
"""
    before_drafts = db_session.query(OnboardingDraft).count()
    before_documents = db_session.query(SourceDocument).count()

    response = client.post(
        "/api/onboarding/uploads/extract",
        headers=headers,
        files=[("files", ("misleading-document-name.txt", content, "text/plain"))],
    )

    assert response.status_code == 200
    extraction = response.json()
    assert extraction["account_name"] == "Cafe Zupas"
    assert extraction["project_name"] == "Guest Experience Modernization"
    assert extraction["company_url"] == "https://www.cafezupas.com"
    assert extraction["linkedin_url"] == "https://www.linkedin.com/company/cafe-zupas"
    assert extraction["source_file_names"] == ["misleading-document-name.txt"]
    assert db_session.query(OnboardingDraft).count() == before_drafts
    assert db_session.query(SourceDocument).count() == before_documents


def test_onboarding_upload_extract_endpoint_maps_excel_sow_title_and_key_value_rows(client: TestClient) -> None:
    from openpyxl import Workbook

    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SOW"
    for row in (
        ("Cafe Zupas - Statement of Work",),
        ("Project validation baseline prepared from tkxel delivery team",),
        ("Company URL", "https://www.cafezupas.com"),
        ("LinkedIn URL", "https://www.linkedin.com/company/cafe-zupas"),
        ("Service Lines", "Product Engineering, QA Automation, API Integration"),
        ("Contract Value", "USD 240,000"),
        ("Start Date", "01/15/2026"),
        ("End Date", "12/31/2026"),
    ):
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)

    try:
        response = client.post(
            "/api/onboarding/uploads/extract",
            headers=headers,
            files=[
                (
                    "files",
                    (
                        "Cafe_Zupas_SOW.xlsx",
                        stream.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                )
            ],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert response.status_code == 200
    extraction = response.json()
    assert extraction["account_name"] == "Cafe Zupas"
    assert extraction["project_name"] == "Project validation baseline"
    assert extraction["company_url"] == "https://www.cafezupas.com"
    assert extraction["linkedin_url"] == "https://www.linkedin.com/company/cafe-zupas"
    assert extraction["source_file_names"] == ["Cafe_Zupas_SOW.xlsx"]


def test_create_account_charter_upload_uses_excel_parser_draft_edits_and_am_approval(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from openpyxl import Workbook

    def fail_ai_call(*args, **kwargs):
        raise AssertionError("Create Account charter upload must not call AI extraction when use_ai=false")

    monkeypatch.setattr(SowExtractionService, "_call_openai", fail_ai_call)
    monkeypatch.setattr(SowExtractionService, "_call_local_ai", fail_ai_call)
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = True
    admin_headers = auth_headers(client)
    owner = seeded_user(client, admin_headers, "account_manager")
    kam_head = seeded_user(client, admin_headers, "kam_head")
    owner_headers = auth_headers(client, owner["email"], "User@12345")
    kam_headers = auth_headers(client, kam_head["email"], "User@12345")

    workbook = Workbook()
    project_info = workbook.active
    project_info.title = "Project Info"
    for row in (
        ("Account Name", "ASAP Semiconductor"),
        ("Account Executive", "Kamran"),
        ("Division - Department", "Professional Services - Engineering"),
        ("Project Kickoff Date", 46035),
        ("Project Name", "ASAP - Bespoke CMS Project - SOW#02"),
        ("Contract Type", "Software Development"),
        ("Resource Agreement Type", "Fixed Price"),
        ("Project Size (man hours)", 5600),
    ):
        project_info.append(row)
    scope = workbook.create_sheet("Scope")
    for row in (
        ("Industry Vertical", "Airline"),
        ("Business Domain", "ASAP needs a bespoke CMS to manage product data and vendor workflows."),
        ("Project Domain", "Content Management"),
        ("Project Objectives", "Build a searchable, governed CMS for internal teams and vendors."),
        ("Invoicing Methodology", "Advance"),
        ("Invoicing Schedule", "Shared with client spread across 36 months with Net-30 days payment terms"),
    ):
        scope.append(row)
    streams = workbook.create_sheet("Streams & Compliance")
    streams.append(("Service Stream Name", "CMS Development - Software Development"))
    streams.append(("JIRA Project Name", "ASAP CMS"))
    risks = workbook.create_sheet("Risks")
    risks.append(("Risk Statement", "Mitigation Plan"))
    risks.append(("Data Quality Issues (Duplicates / Incorrect Product Data)", "Run validation rules before migration"))
    risks.append(("Security / Unauthorized Vendor Access", "Add role based access controls"))
    stream = BytesIO()
    workbook.save(stream)
    charter_file = (
        "ASAP_Charter.xlsx",
        stream.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    try:
        preview = client.post(
            "/api/onboarding/uploads/extract",
            headers=admin_headers,
            data={"use_ai": "false"},
            files=[("files", charter_file)],
        )
        assert preview.status_code == 200
        assert preview.json()["account_name"] == "ASAP Semiconductor"
        assert preview.json()["project_name"] == "ASAP - Bespoke CMS Project - SOW#02"

        stream.seek(0)
        created = client.post(
            "/api/onboarding/drafts/upload",
            headers=admin_headers,
            data={"manager_id": owner["id"], "use_ai": "false"},
            files=[("files", charter_file)],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert created.status_code == 201
    draft = created.json()
    assert draft["account_name"] == "ASAP Semiconductor"
    assert draft["project_name"] == "ASAP - Bespoke CMS Project - SOW#02"
    assert draft["company_url"] is None
    assert draft["linkedin_url"] is None
    assert draft["segment"] == "Airline"
    assert "Business Domain: ASAP needs a bespoke CMS" in draft["service_context"]
    assert "Project Size (man hours): 5600" in draft["commercial_summary"]
    assert "Data Quality Issues" in draft["initial_notes"]
    assert draft["engagement_drafts"][0]["service_lines"] == ["CMS Development - Software Development"]
    assert draft["engagement_drafts"][0]["start_date"].startswith("2026-01-13")

    owner_notifications = client.get("/api/notifications", headers=owner_headers, params={"trigger": "account_draft_created"})
    assert owner_notifications.status_code == 200
    assert any(item["source_record_id"] == draft["id"] for item in owner_notifications.json()["items"])
    kam_created_notifications = client.get("/api/notifications", headers=kam_headers, params={"trigger": "account_draft_created"})
    assert kam_created_notifications.status_code == 200
    assert any(item["source_record_id"] == draft["id"] for item in kam_created_notifications.json()["items"])

    edit_response = client.patch(
        f"/api/onboarding/drafts/{draft['id']}",
        headers=owner_headers,
        json={
            "company_url": "https://www.asapsemi.com",
            "linkedin_url": "https://www.linkedin.com/company/asap-semiconductor",
            "primary_owner_id": owner["id"],
        },
    )
    assert edit_response.status_code == 200
    assert edit_response.json()["company_url"] == "https://www.asapsemi.com"
    kam_updated_notifications = client.get("/api/notifications", headers=kam_headers, params={"trigger": "account_draft_updated"})
    assert kam_updated_notifications.status_code == 200
    assert any(
        item["source_record_id"] == draft["id"]
        and item["title"] == "Draft updated: ASAP Semiconductor"
        and "company URL" in item["body"]
        and "LinkedIn URL" in item["body"]
        for item in kam_updated_notifications.json()["items"]
    )
    owner_updated_notifications = client.get("/api/notifications", headers=owner_headers, params={"trigger": "account_draft_updated"})
    assert owner_updated_notifications.status_code == 200
    assert not any(item["source_record_id"] == draft["id"] for item in owner_updated_notifications.json()["items"])

    approved = client.post(f"/api/onboarding/drafts/{draft['id']}/approve", headers=owner_headers)
    assert approved.status_code == 200
    assert approved.json()["status"] == "approved"
    assert approved.json()["approved_account_id"]

    kam_approved_notifications = client.get("/api/notifications", headers=kam_headers, params={"trigger": "account_draft_approved"})
    assert kam_approved_notifications.status_code == 200
    assert any(
        item["source_record_id"] == approved.json()["approved_account_id"]
        and item["title"] == "New account onboarded: ASAP Semiconductor"
        for item in kam_approved_notifications.json()["items"]
    )


def test_onboarding_upload_leaves_unknown_fields_blank_instead_of_using_filename(client: TestClient) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    content = b"""
Meeting notes and operational context.
The source document mentions kickoff logistics and discovery workshops but does not include a client legal name,
project title, company website, or LinkedIn URL.
"""

    try:
        response = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            files=[("files", ("cafe-zupas-account-sow.txt", content, "text/plain"))],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert response.status_code == 201
    draft = response.json()
    assert draft["account_name"] == ""
    assert draft["project_name"] == ""
    assert draft["company_url"] is None
    assert draft["linkedin_url"] is None
    assert any("Account name was not found" in field for field in draft["missing_fields"])


def test_onboarding_draft_flags_duplicate_company_url_before_approval(client: TestClient) -> None:
    headers = auth_headers(client)
    account_id, owner_id, approved = create_approved_account(client, headers, "Duplicate Company URL Base")
    duplicate_url = approved["company_url"]

    draft_response = client.post(
        "/api/onboarding/drafts",
        headers=headers,
        json={**draft_payload("Different Name Same Website", owner_id), "company_url": duplicate_url},
    )

    assert draft_response.status_code == 201
    draft = draft_response.json()
    assert draft["duplicate_account_id"] == account_id
    assert any("Possible duplicate account" in conflict for conflict in draft["conflicts"])

    approve_response = client.post(f"/api/onboarding/drafts/{draft['id']}/approve", headers=headers)
    assert approve_response.status_code == 409


def test_onboarding_upload_allows_reusing_same_source_file(client: TestClient, db_session: Session) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    existing_drafts = db_session.query(OnboardingDraft).count()
    existing_documents = db_session.query(SourceDocument).count()
    content = b"""
Account Name: Duplicate Source Customer
Project Name: Duplicate SOW Test
Service Lines: Product Engineering
Contract Value: USD 10000
Start Date: 2026-07-01
End Date: 2026-12-31
"""

    try:
        first = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            files=[("files", ("duplicate-sow.txt", content, "text/plain"))],
        )
        assert first.status_code == 201

        second = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            files=[("files", ("duplicate-sow-copy.txt", content, "text/plain"))],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert second.status_code == 201
    assert first.json()["source_documents"][0]["checksum_sha256"] == second.json()["source_documents"][0]["checksum_sha256"]
    assert db_session.query(OnboardingDraft).count() == existing_drafts + 2
    assert db_session.query(SourceDocument).count() == existing_documents + 2


def test_onboarding_source_extraction_retry_endpoint_reextracts_document(client: TestClient) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    content = b"""
Account Name: Retry Source Customer
Project Name: Retry SOW Test
Service Lines: Product Engineering
Contract Value: USD 12000
Start Date: 2026-07-01
End Date: 2026-12-31
"""

    try:
        draft_response = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            files=[("files", ("retry-sow.txt", content, "text/plain"))],
        )
        assert draft_response.status_code == 201
        draft = draft_response.json()
        document_id = draft["source_documents"][0]["id"]

        retry_response = client.post(
            f"/api/onboarding/drafts/{draft['id']}/documents/{document_id}/extract",
            headers=headers,
            params={"force": "true"},
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert retry_response.status_code == 200
    extraction = retry_response.json()
    assert extraction["source_document_id"] == document_id
    assert extraction["status"] == "completed"
    assert extraction["page_count"] == 1


def test_onboarding_draft_source_reupload_replaces_document_and_refreshes_fields(client: TestClient, db_session: Session) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    wrong_content = b"""
Account Name: Wrong Source Customer
Project Name: Wrong SOW Test
Service Lines: Product Engineering
Contract Value: USD 12000
Start Date: 2026-07-01
End Date: 2026-12-31
"""
    correct_content = b"""
Account Name: Correct Source Customer
Project Name: Correct Charter Replacement
Service Lines: Cloud Migration
Contract Value: USD 45000
Start Date: 2026-08-01
End Date: 2027-02-28
"""

    try:
        draft_response = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            files=[("files", ("wrong-sow.txt", wrong_content, "text/plain"))],
        )
        assert draft_response.status_code == 201
        draft = draft_response.json()
        old_document_id = draft["source_documents"][0]["id"]

        replace_response = client.post(
            f"/api/onboarding/drafts/{draft['id']}/documents/upload",
            headers=headers,
            files=[("files", ("correct-charter.txt", correct_content, "text/plain"))],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert replace_response.status_code == 200
    refreshed = replace_response.json()
    assert refreshed["account_name"] == "Correct Source Customer"
    assert refreshed["project_name"] == "Correct Charter Replacement"
    assert len(refreshed["source_documents"]) == 1
    assert refreshed["source_documents"][0]["file_name"] == "correct-charter.txt"
    assert refreshed["source_documents"][0]["id"] != old_document_id
    assert "Correct Source Customer" in refreshed["source_documents"][0]["extracted_text"]
    assert refreshed["engagement_drafts"][0]["name"] == "Correct Charter Replacement"
    assert db_session.get(SourceDocument, old_document_id) is None
    assert db_session.query(SourceDocument).filter(SourceDocument.draft_id == draft["id"]).count() == 1


def test_pdf_without_readable_text_is_marked_ocr_required(
    db_session: Session,
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_path = tmp_path / "scanned-sow.pdf"
    pdf_path.write_bytes(b"%PDF-1.4 scanned placeholder")
    document = SourceDocument(
        title="Scanned SOW",
        source_type="sow",
        file_name="scanned-sow.pdf",
        storage_path=str(pdf_path),
        mime_type="application/pdf",
        uploaded_by_name="KAM Super Admin",
        extraction_status="queued",
        confidence=0,
        pages=0,
    )
    db_session.add(document)
    db_session.flush()
    service = KycDocumentExtractionService(db_session)
    monkeypatch.setattr(service, "_extract_pdf_pages", lambda path: ([""], "pymupdf"))
    monkeypatch.setattr(service, "_ocr_pdf_pages", lambda path: None)

    extraction = service.extract_document(document, force=True)

    assert extraction.status == "ocr_required"
    assert extraction.error_message == "OCR is required before readable text can be extracted from this source document."
    assert document.extraction_status == "ocr_required"
    assert document.ocr_status == "ocr_required"
    assert document.extraction_error == extraction.error_message
    assert extraction.metadata_json["ocr_required"] is True
    assert db_session.query(DocumentExtraction).filter(DocumentExtraction.document_id == document.id).count() == 0


def test_onboarding_upload_extracts_customer_from_contract_style_sow(client: TestClient, db_session: Session) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    content = b"""
STATEMENT OF WORK - 12
DEDICATED DIGITAL EXPERIENCE TEAM
This Statement of Work ("SOW_MCD_TKLLC_012") adopts and incorporates by reference the terms and
conditions of the Master Services Agreement by and between McDonald's Corporation ("Customer", which term
shall include successors and permitted assigns), with its principal place of business at Chicago, Illinois,
United States, and TkXel LLC ("Consultant"), with its principal place of business at Reston, VA, USA.
1. Scope of Work
1.1 The Consultant shall provide software development services on a full-time dedicated team basis.
Resources
Sr Engineers (Python, React, Java)
QA (functional)
QA Automation
Sr DevOps Engineer (timezone overlap)
Solution Architect
AI/ML SME
UI/UX
Total monthly fee: USD 185,000
Applicable monthly fee: USD 170,000
2. Timeframe and Payment Schedule
2.1 The start date of engagement is 07/01/2026 till 06/30/2027.
2.1.1 This SOW shall stand renewed automatically unless written notice of non-renewal is given by either Party
at least eight (8) weeks prior to expiry of the Term.
"""

    try:
        response = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            files=[("files", ("generic-dedicated-team.txt", content, "text/plain"))],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert response.status_code == 201
    draft = response.json()
    assert draft["account_name"] == "McDonald's Corporation"
    assert draft["project_name"] == "STATEMENT OF WORK - 12 - DEDICATED DIGITAL EXPERIENCE TEAM"
    assert draft["region"] == "United States"
    assert draft["commercial_value"] == 185000
    engagement = draft["engagement_drafts"][0]
    assert "Python engineering" in engagement["service_lines"]
    assert "QA automation" in engagement["service_lines"]
    assert engagement["start_date"].startswith("2026-07-01")
    assert engagement["end_date"].startswith("2027-06-30")
    assert engagement["notice_period_days"] == 56
    field_keys = {citation["field_key"]: citation for citation in draft["source_documents"][0]["citations"]}
    assert "account_name" in field_keys
    assert "commercial_value" in field_keys
    assert "start_date" in field_keys
    assert "end_date" in field_keys
    assert "renewal_terms" in field_keys
    assert "notice_period_days" in field_keys
    assert "service_lines" in field_keys
    extraction = (
        db_session.query(SourceDocumentExtraction)
        .filter(SourceDocumentExtraction.source_document_id == draft["source_documents"][0]["id"])
        .one()
    )
    fields = extraction.metadata_json["sow_structured_extraction"]["fields"]
    assert fields["client_name"]["value"] == "McDonald's Corporation"
    assert fields["client_name"]["citation"]["document"] == "generic-dedicated-team.txt"
    assert fields["client_name"]["citation"]["page"] == 1
    assert fields["client_name"]["missing_evidence"] is None
    page_rows = (
        db_session.query(DocumentExtraction)
        .filter(DocumentExtraction.document_id == draft["source_documents"][0]["id"])
        .order_by(DocumentExtraction.page_number)
        .all()
    )
    assert page_rows
    assert page_rows[0].source_file == "generic-dedicated-team.txt"
    assert page_rows[0].checksum
    assert "McDonald's Corporation" in page_rows[0].raw_text


def test_onboarding_upload_uses_qwen_structured_fields_to_prefill_account_and_engagement(
    client: TestClient,
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = True
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")
    content = b"""
Statement of Work for McDonald's Corporation
Customer: McDonald's Corporation
Program Name: Restaurant Intelligence Data Platform
Company URL: https://www.mcdonalds.com
Region: North America
Service Lines: Data Engineering, Cloud Integration, QA Automation
Contract Value: USD 425000
Start Date: 2030-08-01
End Date: 2031-07-31
Renewal Terms: Automatically renews for one additional 12-month term unless either party gives notice.
Notice Period: 90 days before expiry.
Deliverables:
Unified restaurant analytics, data pipelines, platform QA automation, and executive reporting.
Risks:
Point-of-sale data dependency and franchise rollout sequencing.
"""

    def fake_qwen_response(self: SowExtractionService, prompt: str) -> str:
        assert "You are an SOW extraction engine. Return JSON only." in prompt
        return json.dumps(
            {
                "client_name": {"value": "McDonald's Corporation", "confidence": 0.93},
                "start_date": {"value": "2030-08-01", "confidence": 0.91},
                "end_date": {"value": "2031-07-31", "confidence": 0.91},
                "renewal_terms": {"value": "Automatically renews for one additional 12-month term unless either party gives notice.", "confidence": 0.88},
                "notice_period": {"value": "90 days before expiry", "confidence": 0.89},
                "commercial_value": {"value": 425000, "currency": "USD", "confidence": 0.9},
                "service_lines": {"value": ["Data Engineering", "Cloud Integration", "QA Automation"], "confidence": 0.9},
                "stakeholders": {"value": [], "confidence": 0},
                "deliverables": {"value": ["Unified restaurant analytics", "data pipelines", "platform QA automation"], "confidence": 0.86},
                "risks": {"value": ["Point-of-sale data dependency", "franchise rollout sequencing"], "confidence": 0.84},
            }
        )

    monkeypatch.setattr(SowExtractionService, "_call_local_ai", fake_qwen_response)
    try:
        response = client.post(
            "/api/onboarding/drafts/upload",
            headers=headers,
            data={"manager_id": owner["id"], "manager_name": owner["full_name"], "manager_email": owner["email"]},
            files=[("files", ("not-the-client-name.txt", content, "text/plain"))],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert response.status_code == 201
    draft = response.json()
    assert draft["account_name"] == "McDonald's Corporation"
    assert draft["account_name"] != "Not The Client Name"
    assert draft["project_name"] == "Restaurant Intelligence Data Platform"
    assert draft["company_url"] == "https://www.mcdonalds.com"
    assert draft["commercial_value"] == 425000
    assert draft["currency"] == "USD"
    assert draft["confidence"] >= 90
    assert "Engagement start date was not found" not in " ".join(draft["missing_fields"])
    engagement = draft["engagement_drafts"][0]
    assert engagement["name"] == "Restaurant Intelligence Data Platform"
    assert engagement["delivery_status"] == "planned"
    assert engagement["start_date"].startswith("2030-08-01")
    assert engagement["end_date"].startswith("2031-07-31")
    assert engagement["renewal_date"].startswith("2031-07-31")
    assert engagement["notice_deadline"].startswith("2031-05-02")
    assert engagement["notice_period_days"] == 90
    assert engagement["auto_renewal"] is True
    assert engagement["value"] == 425000
    assert "Data Engineering" in engagement["service_lines"]
    assert "Unified restaurant analytics" in engagement["commercial_context"]
    assert "Renewal terms: Automatically renews" in engagement["commercial_context"]
    assert engagement["resource_dependency"] == "Point-of-sale data dependency"
    assert "Point-of-sale data dependency" in engagement["risks"]
    assert engagement["source_citation"].startswith("not-the-client-name.txt p1: Renewal terms inferred")
    assert "Automatically renews" in engagement["source_citation"]
    assert engagement["confidence"] >= 88
    field_keys = {citation["field_key"]: citation for citation in draft["source_documents"][0]["citations"]}
    assert field_keys["account_name"]["page_number"] == 1
    assert field_keys["account_name"]["confidence"] == 93
    assert field_keys["commercial_value"]["page_number"] == 1
    assert field_keys["commercial_value"]["confidence"] == 90
    assert field_keys["renewal_terms"]["page_number"] == 1
    assert field_keys["renewal_terms"]["confidence"] == 88
    assert "Automatically renews" in field_keys["renewal_terms"]["excerpt"]
    extraction = (
        db_session.query(SourceDocumentExtraction)
        .filter(SourceDocumentExtraction.source_document_id == draft["source_documents"][0]["id"])
        .one()
    )
    stored = extraction.metadata_json["sow_structured_extraction"]
    assert stored["provider"] == "qwen"
    assert stored["fields"]["client_name"]["confidence"] == 93
    assert stored["fields"]["client_name"]["citation"]["document"] == "not-the-client-name.txt"

    approve_response = client.post(f"/api/onboarding/drafts/{draft['id']}/approve", headers=headers)
    assert approve_response.status_code == 200
    account_id = approve_response.json()["approved_account_id"]
    created_engagement = db_session.query(Engagement).filter(Engagement.account_id == account_id).one()
    linked_document = db_session.query(SourceDocument).filter(SourceDocument.id == draft["source_documents"][0]["id"]).one()
    assert linked_document.account_id == account_id
    assert linked_document.engagement_id == created_engagement.id


def test_account_attachment_upload_stores_page_rows_and_sow_structured_fields(client: TestClient, db_session: Session) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = False
    headers = auth_headers(client)
    account_id, _, _ = create_approved_account(client, headers, "Attachment Target Workspace")
    content = b"""
STATEMENT OF WORK - 12
DEDICATED DIGITAL EXPERIENCE TEAM
This Statement of Work adopts and incorporates by reference the terms and conditions of the Master Services
Agreement by and between McDonald's Corporation ("Customer") and TkXel LLC ("Consultant").
1. Scope of Work
The Consultant shall provide software development services on a full-time dedicated team basis.
Resources
Sr Engineers (Python, React, Java)
QA Automation
Sr DevOps Engineer
Total monthly fee: USD 185,000
2. Timeframe and Payment Schedule
The start date of engagement is 07/01/2026 till 06/30/2027.
This SOW shall stand renewed automatically unless written notice of non-renewal is given by either Party
at least eight (8) weeks prior to expiry of the Term.
"""

    try:
        response = client.post(
            f"/api/accounts/{account_id}/attachments/upload",
            headers=headers,
            data={"title": "Uploaded SOW", "source_type": "sow", "extract_now": "true"},
            files=[("file", ("wrong-name.txt", content, "text/plain"))],
        )
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    assert response.status_code == 201
    document = response.json()
    assert document["title"] == "Uploaded SOW"
    assert document["file_name"] == "wrong-name.txt"
    assert document["extraction_status"] == "completed"

    extraction = (
        db_session.query(SourceDocumentExtraction)
        .filter(SourceDocumentExtraction.source_document_id == document["id"])
        .one()
    )
    fields = extraction.metadata_json["sow_structured_extraction"]["fields"]
    assert fields["client_name"]["value"] == "McDonald's Corporation"
    assert fields["client_name"]["citation"]["document"] == "wrong-name.txt"
    assert fields["client_name"]["missing_evidence"] is None
    assert fields["commercial_value"]["value"] == 185000
    assert fields["commercial_value"]["citation"]["page"] == 1
    page_rows = db_session.query(DocumentExtraction).filter(DocumentExtraction.document_id == document["id"]).all()
    assert len(page_rows) == 1
    assert page_rows[0].source_file == "wrong-name.txt"
    assert page_rows[0].page_number == 1
    assert "McDonald's Corporation" in page_rows[0].raw_text


def test_account_attachment_upload_allows_reusing_same_source_file(client: TestClient, db_session: Session) -> None:
    headers = auth_headers(client)
    account_id, _, _ = create_approved_account(client, headers, "Duplicate Attachment Workspace")
    existing_count = db_session.query(SourceDocument).filter(SourceDocument.account_id == account_id).count()
    content = b"""
Statement of Work
Customer: Duplicate Attachment Workspace
Scope of Work:
Tkxel will provide product engineering and QA automation services.
"""
    first = client.post(
        f"/api/accounts/{account_id}/attachments/upload",
        headers=headers,
        data={"title": "First SOW", "source_type": "sow", "extract_now": "false"},
        files=[("file", ("duplicate-account-sow.txt", content, "text/plain"))],
    )
    assert first.status_code == 201

    second = client.post(
        f"/api/accounts/{account_id}/attachments/upload",
        headers=headers,
        data={"title": "Second SOW", "source_type": "sow", "extract_now": "false"},
        files=[("file", ("duplicate-account-sow-copy.txt", content, "text/plain"))],
    )

    assert second.status_code == 201
    assert first.json()["checksum_sha256"] == second.json()["checksum_sha256"]
    assert db_session.query(SourceDocument).filter(SourceDocument.account_id == account_id).count() == existing_count + 2


def test_sow_structured_qwen_result_is_cited_and_vendor_name_guarded(db_session: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    settings = get_settings()
    previous_ai_setting = settings.sow_ai_extraction_enabled
    settings.sow_ai_extraction_enabled = True
    admin = db_session.query(Account).first()
    created_by_id = admin.created_by_id if admin else None
    document = SourceDocument(
        title="Narrative SOW",
        source_type="sow",
        file_name="narrative-sow.pdf",
        uploaded_by_name="KAM Super Admin",
        uploaded_by_id=created_by_id,
        extraction_status="completed",
        confidence=90,
        pages=2,
    )
    db_session.add(document)
    db_session.flush()
    extraction = SourceDocumentExtraction(
        source_document_id=document.id,
        status="completed",
        extractor_name="local-document-extractor",
        extractor_version="v1",
        mime_type="application/pdf",
        page_count=2,
        raw_text="",
        normalized_text="",
        metadata_json={},
    )
    db_session.add(extraction)
    db_session.flush()
    page_one = DocumentExtraction(
        document_id=document.id,
        extraction_id=extraction.id,
        raw_text=(
            "This Statement of Work is by and between McDonald's Corporation (\"Customer\") and TkXel LLC (\"Consultant\"). "
            "The Consultant shall provide software development services. Total monthly fee: USD 185,000."
        ),
        page_number=1,
        source_file="narrative-sow.pdf",
        checksum="page-one",
    )
    page_two = DocumentExtraction(
        document_id=document.id,
        extraction_id=extraction.id,
        raw_text=(
            "The start date of engagement is 07/01/2026 till 06/30/2027. "
            "This SOW shall stand renewed automatically unless written notice of non-renewal is given by either Party "
            "at least eight (8) weeks prior to expiry of the Term."
        ),
        page_number=2,
        source_file="narrative-sow.pdf",
        checksum="page-two",
    )
    db_session.add_all([page_one, page_two])
    db_session.flush()

    def fake_qwen_response(self: SowExtractionService, prompt: str) -> str:
        assert "You are an SOW extraction engine. Return JSON only." in prompt
        assert "It must not be Tkxel" in prompt
        return json.dumps(
            {
                "client_name": {"value": "TKXEL LLC", "confidence": 99},
                "start_date": {
                    "value": "07/01/2026",
                    "confidence": 94,
                    "citation": {
                        "page": 2,
                        "excerpt": "The start date of engagement is 07/01/2026 till 06/30/2027.",
                    },
                },
                "end_date": {"value": "06/30/2027", "confidence": 94},
                "renewal_terms": {"value": "The SOW automatically renews unless either party gives notice.", "confidence": 91},
                "notice_period": {"value": "eight (8) weeks prior to expiry", "confidence": 90},
                "commercial_value": {"value": 185000, "currency": "USD", "confidence": 90},
                "service_lines": {"value": ["Software development"], "confidence": 82},
                "stakeholders": {"value": [], "confidence": 0},
                "deliverables": {"value": ["software development services"], "confidence": 80},
                "risks": {"value": [], "confidence": 0},
            }
        )

    monkeypatch.setattr(SowExtractionService, "_call_local_ai", fake_qwen_response)
    try:
        result = SowExtractionService(db_session).extract_structured_fields(document, extraction)
    finally:
        settings.sow_ai_extraction_enabled = previous_ai_setting

    stored = extraction.metadata_json["sow_structured_extraction"]
    fields = stored["fields"]
    assert result.provider == "qwen"
    assert stored["provider"] == "qwen"
    assert stored["status"] == "complete"
    assert fields["client_name"]["value"] == "McDonald's Corporation"
    assert fields["client_name"]["confidence"] == 86
    assert fields["client_name"]["conflicts"] == ["AI returned Tkxel/vendor/provider as client name and the value was rejected."]
    assert {"field_key": "client_name", "message": "AI returned Tkxel/vendor/provider as client name and the value was rejected."} in stored["conflicts"]
    assert "stakeholders" in stored["missing_fields"]
    assert fields["client_name"]["citation"]["document_id"] == document.id
    assert fields["client_name"]["citation"]["page"] == 1
    assert fields["client_name"]["citation"]["confidence"] == 86
    assert fields["client_name"]["citation"]["field_key"] == "client_name"
    assert fields["client_name"]["citation"]["validation_status"] == "derived_from_extracted_text"
    assert "Customer" in fields["client_name"]["citation"]["excerpt"]
    assert fields["commercial_value"]["value"] == 185000
    assert fields["commercial_value"]["citation"]["page"] == 1
    assert fields["commercial_value"]["citation"]["confidence"] == 90
    assert fields["start_date"]["citation"]["page"] == 2
    assert fields["start_date"]["citation"]["confidence"] == 94
    assert fields["start_date"]["citation"]["validation_status"] == "validated_ai_citation"
    assert fields["renewal_terms"]["citation"]["page"] == 2
    assert fields["stakeholders"]["missing_evidence"] == "Stakeholders were not supported by extracted SOW text."
def test_onboarding_account_manager_candidates_are_active_account_managers(client: TestClient) -> None:
    admin_headers = auth_headers(client)
    kam_headers = auth_headers(client, "abdul.rehman@tkxel.io", "User@12345")
    create_account_manager_user(client, admin_headers, "second.account.manager@tkxel.com", "Second Account Manager")

    response = client.get("/api/onboarding/account-managers", headers=kam_headers)

    assert response.status_code == 200
    candidates = response.json()
    emails = {candidate["email"] for candidate in candidates}
    roles = {candidate["role"] for candidate in candidates}
    assert "account.manager.user@tkxel.com" in emails
    assert "second.account.manager@tkxel.com" in emails
    assert roles == {"account_manager"}


def test_onboarding_approval_requires_explicit_account_manager_assignment(client: TestClient) -> None:
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")
    payload = draft_payload("Unassigned Owner Workspace", owner["id"])
    payload.pop("primary_owner_id")

    create_response = client.post("/api/onboarding/drafts", headers=headers, json=payload)
    assert create_response.status_code == 201

    approve_response = client.post(f"/api/onboarding/drafts/{create_response.json()['id']}/approve", headers=headers)
    assert approve_response.status_code == 422
    error = approve_response.json()["detail"]["errors"][0]
    assert error["field"] == "primary_owner_id"
    assert error["message"] == "Assign an account manager before approving this draft."


def test_onboarding_approval_requires_at_least_one_engagement_draft(client: TestClient) -> None:
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")

    create_response = client.post(
        "/api/onboarding/drafts",
        headers=headers,
        json=draft_payload("No Engagement Draft Workspace", owner["id"], include_engagement=False),
    )
    assert create_response.status_code == 201

    approve_response = client.post(f"/api/onboarding/drafts/{create_response.json()['id']}/approve", headers=headers)
    assert approve_response.status_code == 422
    error = approve_response.json()["detail"]["errors"][0]
    assert error["field"] == "engagement_drafts"
    assert error["message"] == "At least one engagement is required before approval."


def test_onboarding_update_clears_editable_missing_field_blockers(client: TestClient) -> None:
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")
    payload = draft_payload("Editable Missing Fields Workspace", owner["id"])
    payload["company_url"] = None
    payload["missing_fields"] = [
        "Company website was not found in the uploaded source text.",
        "Service lines were not found in the uploaded source text.",
        "Service lines were not supported by extracted SOW text.",
        "Commercial value was not found in the uploaded source text.",
        "Commercial value was not supported by extracted SOW text.",
        "Engagement start date was not found in the uploaded source text.",
        "Start date was not supported by extracted SOW text.",
        "Engagement end date was not found in the uploaded source text.",
        "End date was not supported by extracted SOW text.",
        "Renewal notice period was not found in the uploaded source text.",
        "Notice period was not supported by extracted SOW text.",
    ]
    payload["engagement_drafts"][0]["service_lines"] = []
    payload["engagement_drafts"][0]["value"] = 0
    payload["engagement_drafts"][0]["start_date"] = None
    payload["engagement_drafts"][0]["end_date"] = None
    payload["engagement_drafts"][0]["notice_deadline"] = None
    payload["engagement_drafts"][0]["notice_period_days"] = None

    create_response = client.post("/api/onboarding/drafts", headers=headers, json=payload)
    assert create_response.status_code == 201
    draft = create_response.json()
    assert len(draft["missing_fields"]) == 11

    update_response = client.patch(
        f"/api/onboarding/drafts/{draft['id']}",
        headers=headers,
        json={
            "company_url": "https://editable-missing-fields.example.com",
            "engagement_drafts": [
                {
                    "id": draft["engagement_drafts"][0]["id"],
                    "service_lines": ["Development"],
                    "value": 50000,
                    "start_date": iso_days_from_now(0),
                    "end_date": iso_days_from_now(120),
                    "notice_deadline": iso_days_from_now(90),
                }
            ],
        },
    )

    assert update_response.status_code == 200
    updated = update_response.json()
    assert updated["company_url"] == "https://editable-missing-fields.example.com"
    assert updated["engagement_drafts"][0]["service_lines"] == ["Development"]
    assert updated["engagement_drafts"][0]["value"] == 50000
    assert updated["missing_fields"] == []


def test_onboarding_drafts_are_visible_to_assigned_manager_not_unrelated_managers(client: TestClient) -> None:
    admin_headers = auth_headers(client)
    owner = seeded_user(client, admin_headers, "account_manager")
    other_owner = create_account_manager_user(client, admin_headers, "unrelated.account.manager@tkxel.com", "Unrelated Account Manager")

    draft_response = client.post("/api/onboarding/drafts", headers=admin_headers, json=draft_payload("Assigned Visibility Workspace", owner["id"]))
    assert draft_response.status_code == 201
    draft_id = draft_response.json()["id"]

    assigned_headers = auth_headers(client, owner["email"], "User@12345")
    assigned_list = client.get("/api/onboarding/drafts", headers=assigned_headers, params={"page": 1, "page_size": 10})
    assert assigned_list.status_code == 200
    assert {item["id"] for item in assigned_list.json()["items"]} == {draft_id}

    assigned_read = client.get(f"/api/onboarding/drafts/{draft_id}", headers=assigned_headers)
    assert assigned_read.status_code == 200
    assert assigned_read.json()["primary_owner_id"] == owner["id"]

    unrelated_headers = auth_headers(client, other_owner["email"], "User@12345")
    unrelated_list = client.get("/api/onboarding/drafts", headers=unrelated_headers, params={"page": 1, "page_size": 10})
    assert unrelated_list.status_code == 200
    assert unrelated_list.json()["items"] == []
    assert unrelated_list.json()["total"] == 0

    unrelated_read = client.get(f"/api/onboarding/drafts/{draft_id}", headers=unrelated_headers)
    assert unrelated_read.status_code == 403


def test_account_manager_can_self_assign_but_not_assign_other_managers(client: TestClient) -> None:
    admin_headers = auth_headers(client)
    owner = seeded_user(client, admin_headers, "account_manager")
    other_owner = create_account_manager_user(client, admin_headers, "handoff.account.manager@tkxel.com", "Handoff Account Manager")
    owner_headers = auth_headers(client, owner["email"], "User@12345")

    create_other_response = client.post("/api/onboarding/drafts", headers=owner_headers, json=draft_payload("AM Other Owner Workspace", other_owner["id"]))
    assert create_other_response.status_code == 403

    create_self_response = client.post("/api/onboarding/drafts", headers=owner_headers, json=draft_payload("AM Self Owner Workspace", owner["id"]))
    assert create_self_response.status_code == 201
    draft_id = create_self_response.json()["id"]

    reassign_response = client.patch(f"/api/onboarding/drafts/{draft_id}", headers=owner_headers, json={"primary_owner_id": other_owner["id"]})
    assert reassign_response.status_code == 403


def test_onboarding_validation_and_authorization_errors_are_enforced(client: TestClient) -> None:
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")

    invalid_response = client.post(
        "/api/onboarding/drafts",
        headers=headers,
        json={**draft_payload("Invalid Workspace", owner["id"]), "source_documents": []},
    )
    assert invalid_response.status_code == 422

    invalid_linkedin_response = client.post(
        "/api/onboarding/drafts",
        headers=headers,
        json={**draft_payload("Invalid LinkedIn Workspace", owner["id"]), "linkedin_url": "https://customer.example.com/company"},
    )
    assert invalid_linkedin_response.status_code == 422
    assert invalid_linkedin_response.json()["errors"][0]["field"] == "linkedin_url"
    assert invalid_response.json()["message"] == "Validation failed"

    draft_response = client.post("/api/onboarding/drafts", headers=headers, json=draft_payload("Authorization Workspace", owner["id"]))
    assert draft_response.status_code == 201

    account_manager_headers = auth_headers(client, "account.manager.user@tkxel.com", "User@12345")
    approve_response = client.post(f"/api/onboarding/drafts/{draft_response.json()['id']}/approve", headers=account_manager_headers)
    assert approve_response.status_code == 200
    assert approve_response.json()["approved_account_id"]


def test_engagement_list_endpoint_returns_items_empty_state_and_rejects_unauthorized_access(client: TestClient, db_session: Session) -> None:
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")
    account_id = create_account_without_engagement(db_session, owner, "Engagement List API Workspace")
    owner_id = owner["id"]

    empty_response = client.get(f"/api/accounts/{account_id}/engagements", headers=headers)
    assert empty_response.status_code == 200
    assert empty_response.json()["items"] == []
    assert empty_response.json()["total"] == 0

    created = create_engagement_for_account(client, headers, account_id, owner_id, name="Endpoint list coverage")
    list_response = client.get(f"/api/accounts/{account_id}/engagements", headers=headers, params={"search": "Endpoint list", "page": 1, "page_size": 10})
    assert list_response.status_code == 200
    page = list_response.json()
    assert page["total"] == 1
    assert page["items"][0]["id"] == created["id"]
    assert page["items"][0]["account_id"] == account_id

    unauthorized_user = create_account_manager_user(client, headers, "unassigned.engagement.viewer@tkxel.com", "Unassigned Engagement Viewer")
    unauthorized_headers = auth_headers(client, unauthorized_user["email"], "User@12345")
    unauthorized_response = client.get(f"/api/accounts/{account_id}/engagements", headers=unauthorized_headers)
    assert unauthorized_response.status_code == 403


def test_engagement_list_normalizes_legacy_source_link_routes(client: TestClient, db_session: Session) -> None:
    headers = auth_headers(client)
    account_id, owner_id, _ = create_approved_account(client, headers, "Legacy Engagement Source Link Workspace")
    created = create_engagement_for_account(client, headers, account_id, owner_id, name="Legacy source link coverage")
    engagement = db_session.get(Engagement, created["id"])
    assert engagement is not None
    engagement.source_links = [{"label": "Demo SOW", "source_type": "demo_seed", "route": f"/accounts/{account_id}?tab=kyc"}]
    db_session.commit()

    list_response = client.get(f"/api/accounts/{account_id}/engagements", headers=headers, params={"page": 1, "page_size": 10})

    assert list_response.status_code == 200
    listed = next(item for item in list_response.json()["items"] if item["id"] == created["id"])
    assert listed["source_links"][0]["title"] == "Demo SOW"
    assert listed["source_links"][0]["url"] == f"/accounts/{account_id}?tab=kyc"


def test_engagement_create_endpoint_validates_payload_calculates_notice_and_returns_shape(client: TestClient) -> None:
    headers = auth_headers(client)
    account_id, owner_id, _ = create_approved_account(client, headers, "Engagement Create API Workspace")
    spoofed_owner = create_account_manager_user(client, headers, "spoofed.engagement.owner@tkxel.com", "Spoofed Engagement Owner")

    create_payload = engagement_payload(owner_id)
    create_payload.pop("owner_id")
    create_response = client.post(f"/api/accounts/{account_id}/engagements", headers=headers, json=create_payload)
    assert create_response.status_code == 201
    created = create_response.json()
    assert ENGAGEMENT_RESPONSE_FIELDS.issubset(created.keys())
    assert created["account_id"] == account_id
    assert created["name"] == "Strategic SOW QA"
    assert created["owner_id"] == owner_id
    assert created["contract_value"] == 87500
    assert created["health_score"] == 82
    assert created["source_links"][0]["url"] == "https://customer.example.com/sow"
    assert created["notice_deadline"].startswith(date_days_from_now(45))
    assert created["days_to_expiry"] == 120
    assert created["renewal_status"] == "upcoming_notice_window"

    spoofed_owner_response = client.post(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        json=engagement_payload(spoofed_owner["id"], name="Spoofed owner attempt"),
    )
    assert spoofed_owner_response.status_code == 201
    assert spoofed_owner_response.json()["owner_id"] == owner_id

    missing_name = engagement_payload(owner_id)
    missing_name.pop("name")
    missing_name_response = client.post(f"/api/accounts/{account_id}/engagements", headers=headers, json=missing_name)
    assert missing_name_response.status_code == 422
    assert any(error["field"] == "name" for error in missing_name_response.json()["errors"])

    bad_dates_response = client.post(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        json=engagement_payload(owner_id, start_date=iso_days_from_now(5), end_date=iso_days_from_now(1)),
    )
    assert bad_dates_response.status_code == 422

    negative_value_response = client.post(f"/api/accounts/{account_id}/engagements", headers=headers, json=engagement_payload(owner_id, contract_value=-1))
    assert negative_value_response.status_code == 422

    negative_notice_response = client.post(f"/api/accounts/{account_id}/engagements", headers=headers, json=engagement_payload(owner_id, notice_period_days=-1))
    assert negative_notice_response.status_code == 422


def test_engagement_create_from_charter_upload_creates_editable_draft_notifications_and_approval(
    client: TestClient,
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "sow_ai_extraction_enabled", True)

    def fail_ai(*args, **kwargs):
        raise AssertionError("Engagement charter import must use deterministic document parsing, not AI extraction")

    monkeypatch.setattr(SowExtractionService, "_extract_with_openai", fail_ai)
    monkeypatch.setattr(SowExtractionService, "_extract_with_qwen", fail_ai)

    admin_headers = auth_headers(client)
    owner = seeded_user(client, admin_headers, "account_manager")
    other_owner = create_account_manager_user(client, admin_headers, "charter.other.owner@tkxel.com", "Charter Other Owner")
    kam_head = seeded_user(client, admin_headers, "kam_head")
    owner_headers = auth_headers(client, owner["email"], "User@12345")
    kam_headers = auth_headers(client, kam_head["email"], "User@12345")
    account_id = create_account_without_engagement(db_session, owner, "Charter Import Workspace")
    content = b"""
Project Name: Restaurant Digital Experience Modernization
Client: Cafe Zupas
Start Date: 01/15/2026
End Date: 12/31/2026
Contract Value: USD 240,000
Service Lines: Product Engineering, QA Automation, DevOps
Stakeholders:
Jane Sponsor - VP Digital
Mark Owner - Product Owner
Deliverables: customer ordering platform, analytics dashboard, loyalty integrations
Risks: POS integration dependency, holiday traffic surge
"""

    response = client.post(
        f"/api/accounts/{account_id}/engagements/from-charter",
        headers=owner_headers,
        files={"file": ("cafe-zupas-charter.txt", content, "text/plain")},
    )

    assert response.status_code == 201
    draft = response.json()
    assert draft["account_id"] == account_id
    assert draft["status"] == "ready_for_review"
    assert draft["name"] == "Restaurant Digital Experience Modernization"
    assert draft["contract_value"] == 240000
    assert "product engineering" in {item.lower() for item in draft["service_lines"]}
    assert draft["source_document_ids"]
    assert draft["approved_engagement_id"] is None
    assert db_session.query(Engagement).filter(Engagement.account_id == account_id).count() == 0
    assert db_session.get(EngagementImportDraft, draft["id"]) is not None

    source_document = db_session.query(SourceDocument).filter(SourceDocument.id == draft["source_document_ids"][0]).one()
    assert source_document.account_id == account_id
    assert source_document.engagement_id is None
    assert source_document.extraction_status == "completed"
    assert db_session.query(DocumentExtraction).filter(DocumentExtraction.document_id == source_document.id).count() >= 1

    list_response = client.get(f"/api/accounts/{account_id}/engagement-drafts", headers=owner_headers)
    assert list_response.status_code == 200
    assert list_response.json()["items"][0]["id"] == draft["id"]

    kam_created_notifications = client.get("/api/notifications", headers=kam_headers, params={"trigger": "engagement_draft_created"})
    assert kam_created_notifications.status_code == 200
    assert any(item["source_record_id"] == draft["id"] for item in kam_created_notifications.json()["items"])

    update_response = client.patch(
        f"/api/engagement-drafts/{draft['id']}",
        headers=owner_headers,
        json={
            "name": "Reviewed Restaurant Digital Experience",
            "owner_id": other_owner["id"],
            "contract_value": 260000,
            "risks": ["POS integration dependency"],
        },
    )
    assert update_response.status_code == 200
    updated = update_response.json()
    assert updated["name"] == "Reviewed Restaurant Digital Experience"
    assert updated["owner_id"] == other_owner["id"]
    assert updated["contract_value"] == 260000

    kam_updated_notifications = client.get("/api/notifications", headers=kam_headers, params={"trigger": "engagement_draft_updated"})
    assert kam_updated_notifications.status_code == 200
    assert any(item["source_record_id"] == draft["id"] for item in kam_updated_notifications.json()["items"])

    approve_response = client.post(f"/api/engagement-drafts/{draft['id']}/approve", headers=owner_headers)
    assert approve_response.status_code == 200
    approved = approve_response.json()
    assert approved["status"] == "approved"
    assert approved["approved_engagement_id"]

    created = db_session.get(Engagement, approved["approved_engagement_id"])
    assert created is not None
    assert created.account_id == account_id
    assert created.name == "Reviewed Restaurant Digital Experience"
    assert created.owner_id == owner["id"]
    assert created.owner_name == owner["full_name"]
    assert float(created.value) == 260000
    db_session.refresh(source_document)
    assert source_document.engagement_id == created.id
    stakeholder_names = {item.name for item in db_session.query(Stakeholder).filter(Stakeholder.engagement_id == created.id).all()}
    assert "Jane Sponsor" in stakeholder_names
    assert "Mark Owner" in stakeholder_names

    kam_approved_notifications = client.get("/api/notifications", headers=kam_headers, params={"trigger": "engagement_draft_approved"})
    assert kam_approved_notifications.status_code == 200
    assert any(item["source_record_id"] == created.id for item in kam_approved_notifications.json()["items"])


def test_engagement_detail_endpoint_returns_detail_404_and_rejects_unauthorized_access(client: TestClient) -> None:
    headers = auth_headers(client)
    account_id, owner_id, _ = create_approved_account(client, headers, "Engagement Detail API Workspace")
    created = create_engagement_for_account(client, headers, account_id, owner_id, name="Endpoint detail coverage")

    detail_response = client.get(f"/api/engagements/{created['id']}", headers=headers)
    assert detail_response.status_code == 200
    detail = detail_response.json()
    assert ENGAGEMENT_RESPONSE_FIELDS.issubset(detail.keys())
    assert detail["id"] == created["id"]
    assert detail["account_id"] == account_id
    assert detail["account_name"] == "Engagement Detail API Workspace"
    assert detail["name"] == "Endpoint detail coverage"

    missing_response = client.get("/api/engagements/missing-engagement", headers=headers)
    assert missing_response.status_code == 404

    unauthorized_user = create_account_manager_user(client, headers, "unassigned.engagement.detail@tkxel.com", "Unassigned Engagement Detail")
    unauthorized_headers = auth_headers(client, unauthorized_user["email"], "User@12345")
    unauthorized_response = client.get(f"/api/engagements/{created['id']}", headers=unauthorized_headers)
    assert unauthorized_response.status_code == 403


def test_engagement_patch_endpoint_updates_recalculates_notice_and_validates_payload(client: TestClient) -> None:
    headers = auth_headers(client)
    account_id, owner_id, _ = create_approved_account(client, headers, "Engagement Patch API Workspace")
    spoofed_owner = create_account_manager_user(client, headers, "patch.spoofed.owner@tkxel.com", "Patch Spoofed Owner")
    created = create_engagement_for_account(client, headers, account_id, owner_id, name="Endpoint patch coverage")

    update_response = client.patch(
        f"/api/engagements/{created['id']}",
        headers=headers,
        json={
            "name": "Endpoint patch updated",
            "description": "Updated description",
            "owner_id": spoofed_owner["id"],
            "service_lines": ["Cloud", "Data"],
            "contract_value": 99000,
            "currency": "EUR",
            "delivery_status": "watch",
            "commercial_status": "risk",
            "health_score": 66,
            "health_status": "amber",
            "renewal_risk": "high",
            "end_date": iso_days_from_now(120),
            "renewal_date": iso_days_from_now(40),
            "notice_period_days": 12,
            "notice_deadline": iso_days_from_now(1),
            "resource_dependency_notes": "Updated resource note.",
            "risks": ["Updated risk"],
        },
    )
    assert update_response.status_code == 200
    updated = update_response.json()
    assert ENGAGEMENT_RESPONSE_FIELDS.issubset(updated.keys())
    assert updated["name"] == "Endpoint patch updated"
    assert updated["description"] == "Updated description"
    assert updated["owner_id"] == owner_id
    assert updated["service_lines"] == ["Cloud", "Data"]
    assert updated["contract_value"] == 99000
    assert updated["currency"] == "EUR"
    assert updated["delivery_status"] == "watch"
    assert updated["commercial_status"] == "risk"
    assert updated["health_score"] == 66
    assert updated["health_status"] == "amber"
    assert updated["renewal_risk"] == "high"
    assert updated["resource_dependency_notes"] == "Updated resource note."
    assert updated["risks"] == ["Updated risk"]
    assert updated["notice_deadline"].startswith(date_days_from_now(28))
    assert updated["renewal_status"] == "notice_due"

    bad_dates_response = client.patch(
        f"/api/engagements/{created['id']}",
        headers=headers,
        json={"start_date": iso_days_from_now(10), "end_date": iso_days_from_now(3)},
    )
    assert bad_dates_response.status_code == 422

    negative_value_response = client.patch(f"/api/engagements/{created['id']}", headers=headers, json={"contract_value": -10})
    assert negative_value_response.status_code == 422

    negative_notice_response = client.patch(f"/api/engagements/{created['id']}", headers=headers, json={"notice_period_days": -1})
    assert negative_notice_response.status_code == 422


def test_engagement_archive_endpoint_archives_without_hard_delete_and_hides_from_active_list(client: TestClient, db_session: Session) -> None:
    headers = auth_headers(client)
    account_id, owner_id, _ = create_approved_account(client, headers, "Engagement Archive API Workspace")
    created = create_engagement_for_account(client, headers, account_id, owner_id, name="Endpoint archive coverage")

    archive_response = client.delete(f"/api/engagements/{created['id']}", headers=headers)
    assert archive_response.status_code == 200
    assert archive_response.json()["message"] == "Engagement archived successfully"

    archived = db_session.get(Engagement, created["id"])
    assert archived is not None
    assert archived.status == "archived"
    assert archived.archived_at is not None

    active_list = client.get(f"/api/accounts/{account_id}/engagements", headers=headers, params={"search": "Endpoint archive coverage"})
    assert active_list.status_code == 200
    assert active_list.json()["items"] == []

    detail_response = client.get(f"/api/engagements/{created['id']}", headers=headers)
    assert detail_response.status_code == 404


def test_engagement_summary_endpoint_returns_account_overview_and_engagement_360_fields(client: TestClient) -> None:
    headers = auth_headers(client)
    account_id, owner_id, _ = create_approved_account(client, headers, "Engagement Summary API Workspace")
    created = create_engagement_for_account(client, headers, account_id, owner_id, name="Endpoint summary coverage")

    summary_response = client.get(f"/api/engagements/{created['id']}/summary", headers=headers)
    assert summary_response.status_code == 200
    summary = summary_response.json()
    assert summary["id"] == created["id"]
    assert summary["account_id"] == account_id
    assert summary["name"] == "Endpoint summary coverage"
    assert summary["owner_id"] == owner_id
    assert summary["service_lines"] == ["Engineering", "Customer Success"]
    assert summary["source_links"][0]["url"] == "https://customer.example.com/sow"
    assert summary["contract_value"] == 87500
    assert summary["currency"] == "USD"
    assert summary["delivery_status"] == "active"
    assert summary["commercial_status"] == "healthy"
    assert summary["health_score"] == 82
    assert summary["health_status"] == "green"
    assert summary["renewal_risk"] == "medium"
    assert summary["notice_deadline"].startswith(date_days_from_now(45))
    assert summary["days_to_expiry"] == 120
    assert summary["renewal_status"] == "upcoming_notice_window"
    assert summary["resource_dependency_notes"] == "Named platform engineer availability."
    assert summary["risks"] == ["Dependency on customer data access"]
    assert summary["latest_health"]["overall"] == 82
    assert summary["created_at"]
    assert summary["updated_at"]


def test_manual_engagement_crud_calculations_timeline_and_access(client: TestClient) -> None:
    headers = auth_headers(client)
    account_id, owner_id, _ = create_approved_account(client, headers, "Engagement QA Workspace")

    create_response = client.post(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        json=engagement_payload(owner_id),
    )
    assert create_response.status_code == 201
    created = create_response.json()
    engagement_id = created["id"]
    assert created["account_id"] == account_id
    assert created["name"] == "Strategic SOW QA"
    assert created["contract_value"] == 87500
    assert created["notice_deadline"].startswith(date_days_from_now(45))
    assert created["days_to_expiry"] == 120
    assert created["renewal_status"] == "upcoming_notice_window"

    list_response = client.get(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        params={"search": "Strategic", "status": "active", "owner": owner_id, "page": 1, "page_size": 10},
    )
    assert list_response.status_code == 200
    listed_ids = {item["id"] for item in list_response.json()["items"]}
    assert engagement_id in listed_ids

    detail_response = client.get(f"/api/engagements/{engagement_id}", headers=headers)
    assert detail_response.status_code == 200
    assert detail_response.json()["source_links"][0]["url"] == "https://customer.example.com/sow"

    timeline_response = client.get(f"/api/engagements/{engagement_id}/timeline", headers=headers)
    assert timeline_response.status_code == 200
    assert {event["event_type"] for event in timeline_response.json()["items"]} == {"engagement_created"}
    created_event = timeline_response.json()["items"][0]
    assert created_event["account_id"] == account_id
    assert created_event["engagement_id"] == engagement_id
    assert created_event["source_module"] == "engagements"
    assert created_event["actor_id"]
    assert created_event["new_value"]["notice_deadline"].startswith(date_days_from_now(45))

    viewer_headers = auth_headers(client, "leadership.viewer.user@tkxel.com", "User@12345")
    viewer_list = client.get(f"/api/accounts/{account_id}/engagements", headers=viewer_headers)
    assert viewer_list.status_code == 200
    viewer_update = client.patch(f"/api/engagements/{engagement_id}", headers=viewer_headers, json={"health_status": "red"})
    assert viewer_update.status_code == 403

    update_response = client.patch(
        f"/api/engagements/{engagement_id}",
        headers=headers,
        json={
            "name": "Strategic SOW QA - Renewed",
            "service_lines": ["Engineering", "Data"],
            "contract_value": 91000,
            "currency": "EUR",
            "delivery_status": "watch",
            "health_score": 68,
            "health_status": "amber",
            "renewal_risk": "high",
            "start_date": iso_days_from_now(-10),
            "end_date": iso_days_from_now(150),
            "renewal_date": iso_days_from_now(20),
            "notice_period_days": 10,
        },
    )
    assert update_response.status_code == 200
    updated = update_response.json()
    assert updated["name"] == "Strategic SOW QA - Renewed"
    assert updated["contract_value"] == 91000
    assert updated["health_score"] == 68
    assert updated["health_status"] == "amber"
    assert updated["delivery_status"] == "watch"
    assert updated["notice_deadline"].startswith(date_days_from_now(10))
    assert updated["days_to_expiry"] == 150
    assert updated["renewal_status"] == "renewal_due"

    timeline_after_update = client.get(f"/api/engagements/{engagement_id}/timeline", headers=headers)
    assert timeline_after_update.status_code == 200
    event_types = {event["event_type"] for event in timeline_after_update.json()["items"]}
    assert {
        "engagement_created",
        "engagement_updated",
        "sow_terms_updated",
        "renewal_dates_updated",
        "engagement_health_changed",
        "engagement_delivery_status_changed",
    }.issubset(event_types)

    archive_response = client.delete(f"/api/engagements/{engagement_id}", headers=headers)
    assert archive_response.status_code == 200
    assert archive_response.json()["message"] == "Engagement archived successfully"

    archived_detail = client.get(f"/api/engagements/{engagement_id}", headers=headers)
    assert archived_detail.status_code == 404
    active_list = client.get(f"/api/accounts/{account_id}/engagements", headers=headers, params={"search": "Strategic SOW QA - Renewed"})
    assert active_list.status_code == 200
    assert engagement_id not in {item["id"] for item in active_list.json()["items"]}

    timeline_after_archive = client.get(f"/api/engagements/{engagement_id}/timeline", headers=headers)
    assert timeline_after_archive.status_code == 200
    assert "engagement_archived" in {event["event_type"] for event in timeline_after_archive.json()["items"]}


def test_manual_engagement_create_validation_errors(client: TestClient) -> None:
    headers = auth_headers(client)
    account_id, owner_id, _ = create_approved_account(client, headers, "Engagement Validation Workspace")

    missing_name = engagement_payload(owner_id)
    missing_name.pop("name")
    missing_name_response = client.post(f"/api/accounts/{account_id}/engagements", headers=headers, json=missing_name)
    assert missing_name_response.status_code == 422
    assert any(error["field"] == "name" for error in missing_name_response.json()["errors"])

    bad_dates_response = client.post(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        json=engagement_payload(owner_id, start_date=iso_days_from_now(5), end_date=iso_days_from_now(1)),
    )
    assert bad_dates_response.status_code == 422

    negative_value_response = client.post(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        json=engagement_payload(owner_id, contract_value=-1),
    )
    assert negative_value_response.status_code == 422

    negative_notice_response = client.post(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        json=engagement_payload(owner_id, notice_period_days=-1),
    )
    assert negative_notice_response.status_code == 422

    expired_response = client.post(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        json=engagement_payload(
            owner_id,
            name="Expired SOW QA",
            start_date=iso_days_from_now(-90),
            end_date=iso_days_from_now(-1),
            renewal_date=None,
            notice_period_days=None,
        ),
    )
    assert expired_response.status_code == 201
    assert expired_response.json()["renewal_status"] == "expired"


def test_kam_head_sees_portfolio_and_account_manager_is_limited_to_am_assignments(client: TestClient, db_session: Session) -> None:
    admin_headers = auth_headers(client)
    account_manager = seeded_user(client, admin_headers, "account_manager")
    other_manager = create_account_manager_user(client, admin_headers, "scope.other.manager@tkxel.com", "Scope Other Manager")
    kam_head = seeded_user(client, admin_headers, "kam_head")

    assigned_account_id = create_account_without_engagement(db_session, account_manager, "Assigned Account Visibility")
    other_account_id = create_account_without_engagement(db_session, other_manager, "Other Manager Visibility")
    ops_only_account_id = create_account_without_engagement(db_session, other_manager, "Ops Only Visibility")
    db_session.add(
        AccountOwner(
            account_id=ops_only_account_id,
            user_id=account_manager["id"],
            user_name=account_manager["full_name"],
            user_email=account_manager["email"],
            ownership_role="ops_lead",
            is_primary=False,
            created_by_id=account_manager["id"],
        )
    )

    role = db_session.query(Role).filter(Role.slug == "account_manager").one()
    permission = db_session.query(Permission).filter(Permission.module == "accounts", Permission.action == "view_portfolio").one()
    stale_grant = db_session.query(RolePermission).filter(RolePermission.role_id == role.id, RolePermission.permission_id == permission.id).one_or_none()
    if stale_grant is None:
        db_session.add(RolePermission(role_id=role.id, permission_id=permission.id, allowed=True))
    else:
        stale_grant.allowed = True
    db_session.commit()

    manager_headers = auth_headers(client, account_manager["email"], "User@12345")
    manager_capabilities = client.get("/api/users/me/capabilities", headers=manager_headers)
    assert manager_capabilities.status_code == 200
    assert "accounts:view_portfolio" in manager_capabilities.json()["permission_keys"]
    assert manager_capabilities.json()["can_view_portfolio"] is False

    manager_accounts = client.get("/api/accounts", headers=manager_headers, params={"page": 1, "page_size": 500})
    assert manager_accounts.status_code == 200
    manager_payload = manager_accounts.json()
    manager_account_ids = {item["id"] for item in manager_payload["items"]}
    assert manager_payload["page_size"] == 500
    assert manager_payload["total"] == 1
    assert assigned_account_id in manager_account_ids
    assert other_account_id not in manager_account_ids
    assert ops_only_account_id not in manager_account_ids

    assigned_detail = client.get(f"/api/accounts/{assigned_account_id}", headers=manager_headers)
    assert assigned_detail.status_code == 200
    other_detail = client.get(f"/api/accounts/{other_account_id}", headers=manager_headers)
    assert other_detail.status_code == 403
    ops_only_detail = client.get(f"/api/accounts/{ops_only_account_id}", headers=manager_headers)
    assert ops_only_detail.status_code == 403

    kam_headers = auth_headers(client, kam_head["email"], "User@12345")
    kam_accounts = client.get("/api/accounts", headers=kam_headers, params={"page": 1, "page_size": 500})
    assert kam_accounts.status_code == 200
    kam_account_ids = {item["id"] for item in kam_accounts.json()["items"]}
    assert {assigned_account_id, other_account_id, ops_only_account_id}.issubset(kam_account_ids)

    admin_accounts = client.get("/api/accounts", headers=admin_headers, params={"page": 1, "page_size": 500})
    assert admin_accounts.status_code == 200
    admin_account_ids = {item["id"] for item in admin_accounts.json()["items"]}
    assert {assigned_account_id, other_account_id, ops_only_account_id}.issubset(admin_account_ids)

    leadership_headers = auth_headers(client, "leadership.viewer.user@tkxel.com", "User@12345")
    leadership_accounts = client.get("/api/accounts", headers=leadership_headers, params={"page": 1, "page_size": 500})
    assert leadership_accounts.status_code == 200
    leadership_account_ids = {item["id"] for item in leadership_accounts.json()["items"]}
    assert {assigned_account_id, other_account_id, ops_only_account_id}.issubset(leadership_account_ids)


def test_account_filters_owner_history_engagement_health_and_openapi_docs(client: TestClient) -> None:
    headers = auth_headers(client)
    account_id, owner_id, _ = create_approved_account(client, headers, "Globex Workspace")
    ops_lead = seeded_user(client, headers, "kam_head")

    account_list = client.get(
        "/api/accounts",
        headers=headers,
        params={"search": "Globex", "lifecycle_status": "Onboarding", "segment": "Growth", "risk_status": "warning", "page": 1, "page_size": 1},
    )
    assert account_list.status_code == 200
    assert account_list.json()["items"][0]["id"] == account_id
    assert account_list.json()["items"][0]["health"]["overall"] == 45

    email_search = client.get(
        "/api/accounts",
        headers=headers,
        params={"search": "account.manager.user@tkxel.com", "page": 1, "page_size": 5},
    )
    assert email_search.status_code == 200
    assert any(item["id"] == account_id for item in email_search.json()["items"])

    owner_response = client.post(
        f"/api/accounts/{account_id}/owners",
        headers=headers,
        json={"user_id": ops_lead["id"], "ownership_role": "ops_lead", "rationale": "Delivery governance owner."},
    )
    assert owner_response.status_code == 201

    history_response = client.get(f"/api/accounts/{account_id}/ownership-history", headers=headers, params={"page": 1, "page_size": 10})
    assert history_response.status_code == 200
    assert history_response.json()["total"] >= 2

    bad_engagement = client.post(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        json={
            "name": "Bad dates",
            "owner_id": owner_id,
            "service_lines": ["Engineering"],
            "delivery_status": "active",
            "start_date": "2026-05-30T00:00:00Z",
            "end_date": "2026-05-01T00:00:00Z",
        },
    )
    assert bad_engagement.status_code == 422

    created_engagement = client.post(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        json=engagement_payload(owner_id, name="Customer health workspace"),
    )
    assert created_engagement.status_code == 201

    engagements = client.get(
        f"/api/accounts/{account_id}/engagements",
        headers=headers,
        params={"search": "Customer", "service_line": "Engineering", "page": 1, "page_size": 1},
    )
    assert engagements.status_code == 200
    engagement_id = engagements.json()["items"][0]["id"]

    health_status_update = client.patch(f"/api/engagements/{engagement_id}", headers=headers, json={"health_status": "red"})
    assert health_status_update.status_code == 200
    assert health_status_update.json()["health_status"] == "red"

    rollup_response = client.get(f"/api/accounts/{account_id}/health/rollup", headers=headers)
    assert rollup_response.status_code == 200
    rollup = rollup_response.json()
    assert rollup["metric_version"] == "engagement-health-rollup-adapter-v1"
    assert rollup["overall"] == 45
    assert rollup["rag_status"] == "warning"
    assert rollup["contributions"][0]["health_status"] == "red"
    assert rollup["contributions"][0]["scoring_status"] == "pending_scoring_engine"

    timeline_response = client.get(f"/api/engagements/{engagement_id}/timeline", headers=headers)
    assert timeline_response.status_code == 200
    assert any(event["event_type"] == "engagement_health_changed" for event in timeline_response.json()["items"])

    health_response = client.post(f"/api/engagements/{engagement_id}/health/recalculate", headers=headers)
    assert health_response.status_code == 200
    assert health_response.json()["rag_status"] in {"healthy", "warning", "critical"}

    health_history = client.get(f"/api/engagements/{engagement_id}/health", headers=headers, params={"page": 1, "page_size": 5})
    assert health_history.status_code == 200
    assert health_history.json()["total"] >= 2

    openapi = client.get("/openapi.json")
    assert openapi.status_code == 200
    paths = openapi.json()["paths"]
    assert paths["/api/onboarding/drafts"]["post"]["summary"] == "Create onboarding draft"
    assert paths["/api/accounts/{account_id}/overview"]["get"]["summary"] == "Read account overview"
    assert paths["/api/engagements/{engagement_id}/health/recalculate"]["post"]["summary"] == "Recalculate engagement health"


def test_account_list_supports_server_sorting(client: TestClient) -> None:
    headers = auth_headers(client)
    create_approved_account(client, headers, "Aardvark Workspace")
    create_approved_account(client, headers, "Zenith Workspace")

    name_desc = client.get(
        "/api/accounts",
        headers=headers,
        params={"sort": "name", "direction": "desc", "page": 1, "page_size": 2},
    )
    assert name_desc.status_code == 200
    assert [item["name"] for item in name_desc.json()["items"]] == ["Zenith Workspace", "Aardvark Workspace"]

    value_desc = client.get(
        "/api/accounts",
        headers=headers,
        params={"sort": "commercial_value", "direction": "desc", "page": 1, "page_size": 2},
    )
    assert value_desc.status_code == 200
    assert value_desc.json()["items"][0]["commercial_value"] >= value_desc.json()["items"][1]["commercial_value"]

    owner_sort = client.get(
        "/api/accounts",
        headers=headers,
        params={"sort": "owner_name", "direction": "asc", "page": 1, "page_size": 2},
    )
    assert owner_sort.status_code == 200
    assert owner_sort.json()["items"][0]["primary_owner"]["user_name"] == "Account Manager KAM"


def test_accounts_expose_numeric_display_ids_for_search_and_sort(client: TestClient) -> None:
    headers = auth_headers(client)
    first_account_id, _, _ = create_approved_account(client, headers, "Numeric Reference One")
    second_account_id, _, _ = create_approved_account(client, headers, "Numeric Reference Two")

    first_detail = client.get(f"/api/accounts/{first_account_id}", headers=headers)
    second_detail = client.get(f"/api/accounts/{second_account_id}", headers=headers)
    assert first_detail.status_code == 200
    assert second_detail.status_code == 200
    first_number = first_detail.json()["account_number"]
    second_number = second_detail.json()["account_number"]
    assert isinstance(first_number, int)
    assert isinstance(second_number, int)
    assert second_number == first_number + 1

    search_response = client.get("/api/accounts", headers=headers, params={"search": f"Account #{second_number}", "page": 1, "page_size": 10})
    assert search_response.status_code == 200
    assert [item["id"] for item in search_response.json()["items"]] == [second_account_id]

    sorted_response = client.get("/api/accounts", headers=headers, params={"sort": "account_number", "direction": "desc", "page": 1, "page_size": 10})
    assert sorted_response.status_code == 200
    numbers = [item["account_number"] for item in sorted_response.json()["items"]]
    assert numbers == sorted(numbers, reverse=True)


def test_account_creation_accepts_field_builder_values(client: TestClient, db_session: Session) -> None:
    headers = auth_headers(client)
    owner = seeded_user(client, headers, "account_manager")

    field_response = client.post(
        "/api/admin/custom-fields",
        headers=headers,
        json={
            "module": "accounts",
            "field_key": "customer_tier",
            "label": "Customer Tier",
            "field_type": "single_select",
            "options": ["Gold", "Silver"],
            "is_required": True,
            "show_in_detail": True,
        },
    )
    assert field_response.status_code == 201

    definitions = client.get("/api/accounts/custom-fields", headers=headers)
    assert definitions.status_code == 200
    assert any(item["field_key"] == "customer_tier" for item in definitions.json())

    missing_required = client.post(
        "/api/onboarding/drafts",
        headers=headers,
        json=draft_payload("Required Field Workspace", owner["id"]),
    )
    assert missing_required.status_code == 422
    assert missing_required.json()["detail"]["errors"][0]["field"] == "custom_field_values.customer_tier"

    draft_response = client.post(
        "/api/onboarding/drafts",
        headers=headers,
        json={**draft_payload("Field Builder Workspace", owner["id"]), "custom_field_values": {"customer_tier": "Gold"}},
    )
    assert draft_response.status_code == 201
    approve_response = client.post(f"/api/onboarding/drafts/{draft_response.json()['id']}/approve", headers=headers)
    assert approve_response.status_code == 200
    account_id = approve_response.json()["approved_account_id"]

    values = [value for value in db_session.query(CustomFieldValue).filter(CustomFieldValue.record_id == account_id).all()]
    assert values
    assert values[0].value == "Gold"

    account_response = client.get(f"/api/accounts/{account_id}", headers=headers)
    assert account_response.status_code == 200
    assert account_response.json()["custom_field_values"]["customer_tier"] == "Gold"


def test_csv_import_persists_accounts_in_onboarding_hierarchy(client: TestClient, db_session: Session) -> None:
    headers = auth_headers(client)

    response = client.post(
        "/api/accounts/import-csv",
        headers=headers,
        json={
            "duplicate_mode": "skip",
            "source_file_name": "accounts.csv",
            "rows": [
                {
                    "account_name": "CSV Cafe Zupas",
                    "project_name": "CSV Customer Success Workspace",
                    "company_url": "https://cafezupas.example.com",
                    "industry": "Restaurants",
                    "arr": 1260000,
                    "stage": "Onboarding",
                    "owner_email": "account.manager.user@tkxel.com",
                    "segment": "Enterprise",
                    "region": "North America",
                }
            ],
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["created"] == 1
    assert body["failed"] == 0
    result = body["results"][0]
    assert result["status"] == "created"
    assert result["account"]["name"] == "CSV Cafe Zupas"
    assert result["draft_id"]

    account = db_session.query(Account).filter(Account.id == result["account_id"]).one()
    assert account.name == "CSV Cafe Zupas"
    assert account.created_from_draft_id == result["draft_id"]
    assert account.owners[0].user_email == "account.manager.user@tkxel.com"
    assert db_session.query(Engagement).filter(Engagement.account_id == account.id).count() == 1
    document = db_session.query(SourceDocument).filter(SourceDocument.account_id == account.id).one()
    assert document.draft_id == result["draft_id"]
    assert document.source_type == "manual_import"

    list_response = client.get("/api/accounts", headers=headers, params={"search": "CSV Cafe", "page": 1, "page_size": 10})
    assert list_response.status_code == 200
    assert list_response.json()["items"][0]["id"] == account.id

    openapi = client.get("/openapi.json")
    assert openapi.status_code == 200
    assert openapi.json()["paths"]["/api/accounts/import-csv"]["post"]["summary"] == "Import accounts from CSV"


def test_csv_import_reports_row_errors_and_skips_duplicates(client: TestClient) -> None:
    headers = auth_headers(client)
    create_approved_account(client, headers, "CSV Duplicate Workspace")

    duplicate_response = client.post(
        "/api/accounts/import-csv",
        headers=headers,
        json={"duplicate_mode": "skip", "rows": [{"account_name": "CSV Duplicate Workspace"}]},
    )
    assert duplicate_response.status_code == 200
    assert duplicate_response.json()["skipped"] == 1
    assert duplicate_response.json()["results"][0]["status"] == "skipped"

    invalid_response = client.post(
        "/api/accounts/import-csv",
        headers=headers,
        json={"duplicate_mode": "skip", "rows": [{"account_name": "", "arr": -5}]},
    )
    assert invalid_response.status_code == 200
    body = invalid_response.json()
    assert body["failed"] == 1
    assert body["results"][0]["status"] == "failed"
    assert any(error["field"] == "account_name" for error in body["results"][0]["errors"])
